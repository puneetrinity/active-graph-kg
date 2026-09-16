"""Disposable cross-database catchup proof; full built-process cutover is separate.

This first group uses the real Flow consent writer/admission/acknowledgement,
Memory receiver and both actual census schemas. It is not a substitute for the
required built web + API + two-worker public-apply/search/cutover rehearsal.
No provider, production target or credential discovery is permitted.
"""

from __future__ import annotations

import base64
import hashlib
import io
import json
import os
import signal
import socket
import subprocess
import sys
import time
from pathlib import Path
from uuid import uuid4

import psycopg
import pytest
from psycopg import sql
from psycopg.conninfo import conninfo_to_dict

from activekg.api import candidate_consent as consent
from activekg.candidate_index import catchup
from activekg.candidate_index.admission import IndexConfiguration
from activekg.candidate_index.contracts import IndexTunables
from tests.test_candidate_consent_postgres import _send
from tests.test_candidate_index_postgres import (
    POLICY,
    RUNTIME,
    isolated_coordination,  # noqa: F401 - explicit disposable PostgreSQL fixtures
    target,  # noqa: F401
)

pytestmark = pytest.mark.skipif(
    os.getenv("ACTIVEKG_INDEX_XS_DISPOSABLE") != "1",
    reason="explicit two-database fixture not selected",
)


def _flow_dsn(name: str) -> str:
    value = os.environ[name]
    parsed = conninfo_to_dict(value)
    assert parsed["host"] == "127.0.0.1"
    assert parsed["dbname"].startswith("flow_4d_test_")
    assert parsed["user"].startswith("flow_4d_test_")
    return value


def _node(tmp_path: Path, root: Path, dsn: str, body: str, *, data=None, extra=None):
    # Database-only process network boundary, before the first product import.
    # Auth's two pre-existing housekeeping intervals are unreferenced; no web
    # entrypoint, fake authority predicate or real mail transport is started.
    prelude = r"""
import {createRequire,syncBuiltinESMExports} from 'node:module';
import {pathToFileURL} from 'node:url';
import {join} from 'node:path';
const require=createRequire(join(process.env.TEST_FLOW_APP,'package.json'));
const net=require('node:net'), original=net.Socket.prototype.connect;
let forbidden=0;
net.Socket.prototype.connect=function(...args){
 const a=args[0]; const host=typeof a==='object'?a.host:args[1];
 const port=typeof a==='object'?a.port:a;
 if(host!=='127.0.0.1'||Number(port)!==Number(new URL(process.env.DATABASE_URL).port)){
  forbidden++;throw Error('fixture_network_refused');
 }
 return original.apply(this,args);
};syncBuiltinESMExports();
for(const k of ['log','info','warn','error','debug'])console[k]=()=>{};
const interval=globalThis.setInterval;globalThis.setInterval=(...args)=>{const timer=interval(...args);timer.unref();return timer;};
const load=p=>import(pathToFileURL(join(process.env.TEST_FLOW_APP,p)).href);
const input=JSON.parse(await new Promise(resolve=>{let s='';process.stdin.setEncoding('utf8');process.stdin.on('data',x=>s+=x);process.stdin.on('end',()=>resolve(s));}));
const {pool}=await load('server/db.ts');
let answer=null;
try {
"""
    suffix = r"""
 if(forbidden)throw Error('fixture_network_attempt');
 process.stdout.write(JSON.stringify({ok:true,answer}));
}catch(error){process.stdout.write(JSON.stringify({ok:false,code:/^candidate_[a-z_]+$/.test(error?.code??'')?error.code:'fixture_failed'}));process.exitCode=1;}
finally{await pool.end();}
"""
    path = tmp_path / f"module-{uuid4().hex}.mts"
    path.write_text(prelude + body + suffix)
    path.chmod(0o600)
    env = {
        "PATH": os.environ["PATH"],
        "LANG": "C.UTF-8",
        "NODE_ENV": "test",
        "DATABASE_URL": dsn,
        "DATABASE_SSL": "false",
        "PGPOOL_MAX": "2",
        "EMAIL_AUTOMATION_ENABLED": "false",
        "TEST_FLOW_APP": str(root),
        **(extra or {}),
    }
    result = subprocess.run(
        [str(root / "node_modules/.bin/tsx"), str(path)],
        cwd=root,
        env=env,
        input=json.dumps(data),
        text=True,
        capture_output=True,
        timeout=30,
    )
    try:
        payload = json.loads(result.stdout)
        if not payload.get("ok"):
            error = RuntimeError("flow_catchup_admission_refused")
            error.code = payload.get("code")
            raise error
        assert result.returncode == 0
        return payload["answer"]
    finally:
        path.unlink(missing_ok=True)


@pytest.mark.parametrize(
    "variant", ["live", "withdrawn", "orphan", "changed_account", "stale_flow_privacy"]
)
def test_current_consent_catchup_across_real_databases(target, monkeypatch, tmp_path, variant):  # noqa: F811
    root = Path(os.environ["ACTIVEKG_INDEX_XS_FLOW_ROOT"]).resolve() / "VantaHireWebsite"
    assert (root / "server/candidate-index/catchup.ts").is_file()
    fo, fr, fro = [
        _flow_dsn(name)
        for name in (
            "FLOW_INDEX_TEST_OWNER_URL",
            "FLOW_INDEX_TEST_RUNTIME_URL",
            "FLOW_INDEX_TEST_READONLY_URL",
        )
    ]
    monkeypatch.setenv("CANDIDATE_INDEX_API_ENABLED", "false")
    with psycopg.connect(fo, autocommit=True) as flow:
        ft = flow.execute(
            "SELECT target_id FROM schema_control.identity WHERE singleton=true"
        ).fetchone()[0]
        _node(
            tmp_path,
            root,
            fr,
            """
const {Client}=require('pg');
const connect=async dsn=>{const c=new Client({connectionString:dsn});await c.connect();return c;};
const {provisionRuntimeRole}=await load('server/schema-control/runtimeRole.ts');
await provisionRuntimeRole({migrateUrl:process.env.TEST_OWNER,runtimeUrl:process.env.DATABASE_URL,
 runtimeRole:new URL(process.env.DATABASE_URL).username,expectedTargetId:input,
 connectMigration:()=>connect(process.env.TEST_OWNER),connectRuntime:()=>connect(process.env.DATABASE_URL)});
answer=true;
""",
            data=ft,
            extra={"TEST_OWNER": fo},
        )
        ro = conninfo_to_dict(fro)["user"]
        for schema in ("public", "schema_control"):
            flow.execute(
                sql.SQL(
                    "GRANT USAGE ON SCHEMA {} TO {}; GRANT SELECT ON ALL TABLES IN SCHEMA {} TO {}"
                ).format(
                    sql.Identifier(schema),
                    sql.Identifier(ro),
                    sql.Identifier(schema),
                    sql.Identifier(ro),
                )
            )
        user = flow.execute(
            "INSERT INTO users(username,password,role,email_verified,auth_version) VALUES(%s,'fixture-only','candidate',true,1) RETURNING id",
            (f"catchup-{uuid4()}@fixture.invalid",),
        ).fetchone()[0]
        flow.execute(
            "INSERT INTO candidate_privacy_sync_state(consumer_name,status,last_success_at) VALUES('flow','healthy',clock_timestamp()) "
            "ON CONFLICT(consumer_name) DO UPDATE SET status='healthy',last_success_at=clock_timestamp()"
        )
        body = _node(
            tmp_path,
            root,
            fr,
            """
const c=await load('server/candidate-consent/contracts.ts');
const r=await load('server/candidate-consent/repository.ts');
const {randomUUID}=require('node:crypto');
const event=await r.captureConsent('grant',c.grantRequestSchema.parse({request_id:randomUUID(),expected_version:0,
 purpose:c.CONSENT_PURPOSE,copy_version:1,copy_sha256:c.CONSENT_COPY_SHA256,
 profile:{display_name:'Catchup fixture',headline:'Backend Python engineer',location:'',skills:['Python'],linkedin:null},resume_version_id:null}),
 {userId:input,authVersion:1,passwordVersion:c.sha256('fixture-only'),reauthenticatedAt:Date.now()});
const loaded=await r.loadConsentCommand(event.eventId);
const email=(await pool.query('SELECT username FROM users WHERE id=$1',[input])).rows[0].username;
answer={...loaded.command,idempotency_key:c.consentIdentity(loaded.command).idempotencyKey,
 proof:{verified_email:email,privacy_subject:[{identifier_type:'email',value:email}]}};
""",
            data=user,
        )
        grant = consent.ConsentCommand.model_validate(body)
        receipt = _send(grant)
        assert receipt["outcome"] == "granted"
        _node(
            tmp_path,
            root,
            fr,
            """
const claim=(await pool.query("SELECT * FROM flow_claim_candidate_consent_outbox('fixture',1,30000,$1)",[input.event_id])).rows[0];
if(!claim)throw Error('claim_missing');
await pool.query('SELECT flow_ack_candidate_consent_outbox($1,$2,$3)',[claim.outbox_id,claim.generation,JSON.stringify(input)]);
answer=true;
""",
            data=receipt,
        )
        config = IndexConfiguration(POLICY, IndexTunables())
        mt = str(
            target.execute(
                "SELECT target_id FROM activekg_schema_control.target_identity"
            ).fetchone()[0]
        )
        posture = catchup.Posture(
            flow_target=ft,
            memory_target=mt,
            flow_tree="a" * 40,
            memory_tree="b" * 40,
            policy_sha256=POLICY.digest,
            running=True,
        )

        def admit(user_id):
            _node(
                tmp_path,
                root,
                fro,
                """
const {requireCandidatePrivacyAllowed}=await load('server/candidate-privacy/decision.ts');
await requireCandidatePrivacyAllowed({type:'candidate_user',id:input},{globalUse:true,newGlobalOperation:true});answer=true;
""",
                data=user_id,
            )

        ops = catchup.Operations(
            fro,
            RUNTIME,
            lambda: posture,
            admit,
            lambda _: pytest.fail("profile-only object read"),
            config,
        )
        monkeypatch.setattr(catchup, "configuration", lambda: config)
        # Page after the immediately preceding UUID, so this proof selects its
        # real source only, without depending on fixtures from another test.
        from uuid import UUID

        cursor = str(UUID(int=UUID(grant.source.source_id).int - 1))
        key = b"k" * 32
        plan = catchup.census(ops, key=key, cursor=cursor, max_rows=1, max_provider_attempts=5)
        assert len(plan.entries) == 1 and plan.entries[0].source_id == grant.source.source_id
        if variant == "withdrawn":
            from tests.test_candidate_consent_postgres import _command

            assert (
                _send(_command(subject=grant.subject_id, version=2, action="withdraw"))["outcome"]
                == "withdrawn"
            )
        elif variant == "orphan":
            flow.execute("DELETE FROM users WHERE id=%s", (user,))
        elif variant == "changed_account":
            flow.execute("UPDATE users SET auth_version=2 WHERE id=%s", (user,))
        elif variant == "stale_flow_privacy":
            flow.execute(
                "UPDATE candidate_privacy_sync_state SET last_success_at=clock_timestamp()-interval '1 hour' WHERE consumer_name='flow'"
            )
        expected = {
            "live": "captured",
            "withdrawn": "withdrawn_or_superseded",
            "orphan": "source_orphaned",
            "changed_account": "account_changed",
            "stale_flow_privacy": "privacy_unavailable",
        }[variant]
        assert catchup.execute(
            ops, plan=plan, key=key, journal=tmp_path / "journal", operator_approved=True
        ) == {expected: 1}
        scope = f"candidate_{grant.subject_id}"
        target.execute("SELECT set_config('app.current_tenant_id',%s,false)", (scope,))
        assert target.execute(
            "SELECT count(*) FROM candidate_index_sources WHERE source_id=%s",
            (grant.source.source_id,),
        ).fetchone()[0] == int(variant == "live")


# All substitutions below are disposable IO/model boundaries. No route, writer,
# privacy decision, admission, reservation, parser or worker loop is replaced.
MODEL_BOUNDARY = r"""
import sys,types
import numpy as np
class Tokenizer:
 def encode(self,text,**kw):return text.split()
class Encoder:
 def __init__(self,*a,**kw):self.tokenizer=Tokenizer();self.max_seq_length=256
 def encode(self,texts,**kw):
  out=np.zeros((len(texts),384),dtype=np.float32)
  for i,text in enumerate(texts):
   out[i,0]=1.0
   for j,word in enumerate(('python','backend','engineer','database','systems')):
    out[i,j+1]=float(word in text.lower())
  return out
class Ranker:
 def __init__(self,*a,**kw):pass
 def predict(self,pairs):return np.array([1.0 for _ in pairs])
m=types.ModuleType('sentence_transformers');m.SentenceTransformer=Encoder;m.CrossEncoder=Ranker;m.__version__='synthetic-fixture'
sys.modules['sentence_transformers']=m
"""

MEMORY_PROCESS = r"""
import os,sys,socket,runpy,json,asyncio
from pathlib import Path
ports=set(json.loads(os.environ['FIXTURE_PORTS']))
original=socket.socket.connect
def refused():
 with open(os.environ['FIXTURE_FORBIDDEN'],'ab',buffering=0) as f:f.write(b'network\n')
 raise RuntimeError('fixture_network_refused')
def connect(self,address):
 if not isinstance(address,tuple) or address[0] not in ('127.0.0.1','::1') or address[1] not in ports:return refused()
 return original(self,address)
socket.socket.connect=connect
socket.socket.connect_ex=lambda self,address: (connect(self,address) or 0)
lookup=socket.getaddrinfo
def getaddrinfo(host,*args,**kwargs):
 if host not in ('127.0.0.1','::1',b'127.0.0.1',None):return refused()
 return lookup(host,*args,**kwargs)
socket.getaddrinfo=getaddrinfo
bind=socket.socket.bind
def local_bind(self,address):
 if isinstance(address,tuple) and address[0]=='0.0.0.0':address=('127.0.0.1',*address[1:])
 return bind(self,address)
socket.socket.bind=local_bind
sys.path.insert(0,os.environ['FIXTURE_MEMORY_ROOT'])
exec(Path(os.environ['FIXTURE_MODEL_BOUNDARY']).read_text())
# Keep the real isolated child, its stdin limits, parser, timeout, network
# refusals, vector validation and frozen EmbeddingProvider. Only the underlying
# inference library is synthetic; no child inherits API/database credentials.
spawn=asyncio.create_subprocess_exec
async def create(*args,**kw):
 if len(args)==6 and args[1:3]==('-I','-B') and args[4]=='--child' and args[5] in ('embed','rerank'):
  boot='import runpy,sys;'+repr('')+';exec('+repr(Path(os.environ['FIXTURE_MODEL_BOUNDARY']).read_text())+');runpy.run_path('+repr(args[3])+',run_name="__main__")'
  args=(args[0],'-I','-B','-c',boot,'--child',args[5])
 return await spawn(*args,**kw)
asyncio.create_subprocess_exec=create
import httpx
class FakeStream(httpx.AsyncByteStream):
 def __init__(self,value):self.value=value
 async def __aiter__(self):yield self.value
class Boundary(httpx.AsyncBaseTransport):
 def __init__(self):self.real=httpx.AsyncHTTPTransport()
 async def handle_async_request(self,request):
  if request.url.host=='api.groq.com':
   body=json.loads(request.content)
   with open(os.environ['FIXTURE_MODEL_CALLS'],'ab',buffering=0) as f:f.write(b'extract\n')
   extracted={'current_title':'Backend engineer','skills_raw':['Python'],'skills_normalized':['python'],'confidence':0.95}
   raw=json.dumps({'model':body['model'],'choices':[{'finish_reason':'stop','message':{'role':'assistant','content':json.dumps(extracted)}}]}).encode()
   return httpx.Response(200,headers={'content-type':'application/json','content-length':str(len(raw))},stream=FakeStream(raw))
  return await self.real.handle_async_request(request)
 async def aclose(self):await self.real.aclose()
client=httpx.AsyncClient
class Client(client):
 def __init__(self,*args,**kw):
  if kw.get('transport') is None:kw['transport']=Boundary()
  super().__init__(*args,**kw)
httpx.AsyncClient=Client
if sys.argv[1]=='api':
 import uvicorn
 uvicorn.run('activekg.api.main:app',host='127.0.0.1',port=int(os.environ['FIXTURE_API_PORT']),workers=1)
else:runpy.run_module('activekg.'+sys.argv[1]+'.worker',run_name='__main__')
"""


class ProcessStack:
    """Owned loopback processes against explicit disposable, migrated DSNs.

    The separate CI orders own fresh-schema/role proofs. This rehearsal owns
    runtime/readiness and real HTTP/timer publication, never database bootstrap
    with a production owner or a fabricated authority/receipt.
    """

    def __init__(self, path, memory_owner):
        self.path = path
        path.chmod(0o700)
        self.memory = Path(__file__).resolve().parents[1]
        self.flow = Path(os.environ["ACTIVEKG_INDEX_XS_FLOW_ROOT"]) / "VantaHireWebsite"
        self.fo = _flow_dsn("FLOW_INDEX_TEST_OWNER_URL")
        self.fr = _flow_dsn("FLOW_INDEX_TEST_RUNTIME_URL")
        self.owner = memory_owner
        self.owned_databases = []
        self.owned_roles = []
        self.admin = None
        self.owned_memory_connection = None
        self.processes = []
        self.logs = []
        self.ports = [self.port() for _ in range(5)]
        self.redisport, self.apiport, self.webport, self.extractport, self.embedport = self.ports
        self.api_url = f"http://127.0.0.1:{self.apiport}"
        self.web_url = f"http://127.0.0.1:{self.webport}"
        self.base = {
            "PATH": os.environ["PATH"],
            "LANG": "C.UTF-8",
            "NODE_ENV": "test",
            "PYTHONDONTWRITEBYTECODE": "1",
            "HF_HUB_OFFLINE": "1",
            "TRANSFORMERS_OFFLINE": "1",
            "TOKENIZERS_PARALLELISM": "false",
            "OMP_NUM_THREADS": "1",
            "OPENBLAS_NUM_THREADS": "1",
        }

    @staticmethod
    def port():
        with socket.socket() as s:
            s.bind(("127.0.0.1", 0))
            return s.getsockname()[1]

    def file(self, name, text):
        p = self.path / name
        p.write_text(text)
        p.chmod(0o600)
        return p

    def run(self, command, env, cwd=None, timeout=120):
        log = self.path / f"run-{len(self.logs)}.log"
        self.logs.append(log)
        with log.open("w") as out:
            result = subprocess.run(
                list(map(str, command)),
                env={**self.base, **env},
                cwd=cwd or self.flow,
                stdout=out,
                stderr=subprocess.STDOUT,
                timeout=timeout,
            )
        assert result.returncode == 0, f"fixture command failed: {log.name}"

    def start(self, command, env, cwd=None):
        log = self.path / f"process-{len(self.logs)}.log"
        self.logs.append(log)
        with log.open("w") as out:
            process = subprocess.Popen(
                list(map(str, command)),
                env={**self.base, **env},
                cwd=cwd or self.flow,
                stdout=out,
                stderr=subprocess.STDOUT,
                start_new_session=True,
            )
        self.processes.append(process)
        return process

    def stop(self, process):
        if process.poll() is None:
            os.killpg(process.pid, signal.SIGTERM)
            try:
                process.wait(timeout=15)
            except subprocess.TimeoutExpired:
                os.killpg(process.pid, signal.SIGKILL)
                process.wait(timeout=10)

    def close(self):
        for process in reversed(self.processes):
            self.stop(process)
        assert all(p.poll() is not None for p in self.processes)
        # Restore the ordinary generated storage bundle; source was never edited.
        self.run(["npm", "run", "build:candidate-index-storage"], {})
        if self.owned_databases:
            if self.owned_memory_connection:
                self.owned_memory_connection.close()
            for database in reversed(self.owned_databases):
                self.admin.execute(
                    sql.SQL("DROP DATABASE {} WITH (FORCE)").format(sql.Identifier(database))
                )
            for role in reversed(self.owned_roles):
                self.admin.execute(sql.SQL("DROP ROLE {}").format(sql.Identifier(role)))
        if self.admin:
            self.admin.close()

    def bootstrap(self):
        # Fresh databases AND roles prevent synthetic privacy cursors from older
        # matrices masquerading as a current feed. Administrator authority is
        # confined to create/drop and extensions on these exact fresh targets.
        admin_dsn = os.environ["ACTIVEKG_INDEX_XS_ADMIN_DSN"]
        info = conninfo_to_dict(admin_dsn)
        assert info["host"] == "127.0.0.1" and info["dbname"] == "postgres"
        self.admin = psycopg.connect(admin_dsn, autocommit=True)
        assert self.admin.execute("SHOW data_directory").fetchone()[0].startswith("/tmp/ealana-4d-")
        suffix = uuid4().hex[:12]

        def dsn(role, database, password):
            from psycopg.conninfo import make_conninfo

            return make_conninfo(
                host="127.0.0.1", port=info["port"], dbname=database, user=role, password=password
            )

        definitions = {}
        for system in ("flow", "memory"):
            owner = f"{system}_4d_test_{suffix}_owner_test"
            runtime = f"{system}_4d_test_{suffix}_runtime_test"
            database = f"{system}_4d_test_{suffix}_test"
            for role, password, create in [
                (owner, "owner-fixture", True),
                (runtime, "runtime-fixture", False),
            ]:
                self.admin.execute(
                    sql.SQL(
                        "CREATE ROLE {} LOGIN NOINHERIT NOSUPERUSER NOBYPASSRLS {} PASSWORD {}"
                    ).format(
                        sql.Identifier(role),
                        sql.SQL("CREATEROLE" if create else "NOCREATEROLE"),
                        sql.Literal(password),
                    )
                )
                self.owned_roles.append(role)
            self.admin.execute(
                sql.SQL("GRANT {} TO {} WITH ADMIN TRUE, INHERIT FALSE, SET FALSE").format(
                    sql.Identifier(runtime), sql.Identifier(owner)
                )
            )
            self.admin.execute(
                sql.SQL("CREATE DATABASE {} OWNER {}").format(
                    sql.Identifier(database), sql.Identifier(owner)
                )
            )
            self.owned_databases.append(database)
            with psycopg.connect(dsn(info["user"], database, ""), autocommit=True) as connection:
                connection.execute(
                    "CREATE EXTENSION IF NOT EXISTS pgcrypto; CREATE EXTENSION IF NOT EXISTS vector"
                )
            definitions[system] = (
                dsn(owner, database, "owner-fixture"),
                dsn(runtime, database, "runtime-fixture"),
                runtime,
            )
        from urllib.parse import quote

        def url(value):
            fields = conninfo_to_dict(value)
            return f"postgresql://{fields['user']}:{quote(fields['password'])}@127.0.0.1:{fields['port']}/{fields['dbname']}"

        self.fo, self.fr = map(url, definitions["flow"][:2])
        ft = "a" * 48
        _node(
            self.path,
            self.flow,
            self.fr,
            """
const {Client}=require('pg');const connect=async dsn=>{const c=new Client({connectionString:dsn});await c.connect();return c;};
const {runReleaseMigration}=await load('server/schema-control/runner.ts');
const result=await runReleaseMigration({migrationsDir:join(process.env.TEST_FLOW_APP,'server/schema-migrations'),
 creds:{migrateUrl:process.env.TEST_OWNER,expectedTargetId:input,environment:'development',allowFreshInitialization:true},connect});
if(result.applied.length!==13)throw Error('fixture_ledger');
const {provisionRuntimeRole}=await load('server/schema-control/runtimeRole.ts');
await provisionRuntimeRole({migrateUrl:process.env.TEST_OWNER,runtimeUrl:process.env.DATABASE_URL,runtimeRole:new URL(process.env.DATABASE_URL).username,
 expectedTargetId:input,connectMigration:()=>connect(process.env.TEST_OWNER),connectRuntime:()=>connect(process.env.DATABASE_URL)});answer=true;
""",
            data=ft,
            extra={"TEST_OWNER": self.fo},
        )
        mo, self.mr = map(url, definitions["memory"][:2])
        mrole = definitions["memory"][2]
        mt = str(uuid4())
        env = {
            "ACTIVEKG_MIGRATE_DSN": mo,
            "ACTIVEKG_DSN": self.mr,
            "ACTIVEKG_RUNTIME_ROLE": mrole,
            "ACTIVEKG_RUNTIME_PASSWORD": "runtime-fixture",
            "ACTIVEKG_SCHEMA_TARGET_ID": mt,
            "ACTIVEKG_SCHEMA_ENVIRONMENT": "development",
            "ACTIVEKG_SCHEMA_FRESH_INIT": "1",
            "ACTIVEKG_MIGRATION_APPLY": "1",
            "ACTIVEKG_SCHEMA_SOURCE_COMMIT": "a" * 40,
            "CANDIDATE_PRIVACY_INTAKE_ENABLED": "false",
            "CANDIDATE_PRIVACY_HMAC_ACTIVE_VERSION": "1",
            "CANDIDATE_PRIVACY_HMAC_KEY_V1": base64.b64encode(b"\1" * 32).decode(),
            "CANDIDATE_PRIVACY_FLOW_ISSUER": "vantahire",
            "CANDIDATE_PRIVACY_FLOW_ACTOR_ID": "vantahire-backend",
            "CANDIDATE_PRIVACY_SIGNAL_ISSUER": "signal",
            "CANDIDATE_PRIVACY_SIGNAL_ACTOR_ID": "signal-service",
        }
        self.run([sys.executable, "scripts/init_railway_db.py"], env, self.memory)
        self.owner = psycopg.connect(mo, autocommit=True)
        self.owned_memory_connection = self.owner
        assert (
            self.owner.execute("SELECT count(*) FROM public.schema_migrations").fetchone()[0] == 28
        )

    def wait(self, predicate, seconds=75):
        deadline = time.monotonic() + seconds
        while time.monotonic() < deadline:
            exited = [
                (index, p.returncode)
                for index, p in enumerate(self.processes)
                if p not in getattr(self, "retired", []) and p.poll() is not None
            ]
            assert not exited, f"fixture process exited (process index, exit code): {exited}"
            value = predicate()
            if value:
                return value
            time.sleep(0.25)
        raise AssertionError("fixture progress deadline")

    def restart_web(self, **flags):
        self.stop(self.web)
        self.retired = [*getattr(self, "retired", []), self.web]
        self.fe.update(flags)
        self.web = self.start(["node", "--import", self.flow_fence, self.web_build], self.fe)
        import httpx

        def running():
            try:
                return (
                    httpx.get(
                        self.web_url + "/api/csrf-token", trust_env=False, timeout=2
                    ).status_code
                    == 200
                )
            except httpx.HTTPError:
                return False

        self.wait(running)

    @staticmethod
    def csrf(client):
        token = client.get("/api/csrf-token").json()["token"]
        return {
            "x-csrf-token": token,
            "Cookie": "; ".join(f"{k}={v}" for k, v in client.cookies.items()),
        }

    def post(self, client, path, **kwargs):
        return client.post(path, headers=self.csrf(client), **kwargs)

    def login(self, client, email):
        response = self.post(
            client, "/api/login", json={"username": email, "password": "fixture-password"}
        )
        assert response.status_code == 200

    def seed_user(self, flow, role):
        email = f"{role}-{uuid4().hex}@fixture.invalid"
        salt = "fixture-salt"
        password = (
            hashlib.scrypt(
                b"fixture-password", salt=salt.encode(), n=16384, r=8, p=1, dklen=64
            ).hex()
            + "."
            + salt
        )
        user = flow.execute(
            "INSERT INTO users(username,password,role,email_verified,auth_version) VALUES(%s,%s,%s,true,1) RETURNING id",
            (email, password, role),
        ).fetchone()[0]
        return user, email

    def seed_job(self, flow, *, org=None, recruiter=None):
        if recruiter is None:
            recruiter, email = self.seed_user(flow, "recruiter")
        else:
            email = flow.execute("SELECT username FROM users WHERE id=%s", (recruiter,)).fetchone()[
                0
            ]
        if org is None:
            org = flow.execute(
                "INSERT INTO organizations(name,slug,is_active) VALUES('Index fixture',%s,true) RETURNING id",
                ("fixture-" + uuid4().hex,),
            ).fetchone()[0]
            flow.execute(
                "INSERT INTO organization_members(organization_id,user_id,role,seat_assigned) VALUES(%s,%s,'owner',true)",
                (org, recruiter),
            )
        job = flow.execute(
            "INSERT INTO jobs(organization_id,title,location,type,description,posted_by,is_active,status,slug) VALUES(%s,'Backend engineer','Remote','full-time','Python backend systems',%s,true,'approved',%s) RETURNING id",
            (org, recruiter, "fixture-" + uuid4().hex),
        ).fetchone()[0]
        return org, recruiter, email, job

    def privacy_fresh(self, flow):
        def fresh():
            row = flow.execute(
                "SELECT status,last_success_at > clock_timestamp()-interval '1 minute' FROM candidate_privacy_sync_state WHERE consumer_name='flow'"
            ).fetchone()
            return row and row[0] == "healthy" and row[1]

        self.wait(fresh)

    @staticmethod
    def resume():
        from docx import Document

        text = "Backend engineer. Python database development and distributed systems. " * 4
        document = Document()
        document.add_paragraph(text)
        buf = io.BytesIO()
        document.save(buf)
        return text, buf.getvalue()

    def apply(self, client, job, email, *, resume_id=None, recruiter=False):
        data = {
            "name": "Index fixture",
            "email": email,
            "phone": "2025550147",
            "whatsappConsent": "false",
        }
        if resume_id is not None:
            data["resumeId"] = str(resume_id)
            files = None
        else:
            files = {
                "resume": (
                    "fixture.docx",
                    self.resume()[1],
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            }
        route = "applications/recruiter-add" if recruiter else "apply"
        return self.post(client, f"/api/jobs/{job}/{route}", data=data, files=files)

    def published(self, flow, application):
        def complete():
            row = flow.execute(
                "SELECT d.state,d.error_code,d.accepted_source_id FROM candidate_index_outbox o LEFT JOIN candidate_index_delivery_state d USING(outbox_id) WHERE o.application_id=%s",
                (application,),
            ).fetchone()
            self.file("delivery.json", json.dumps(row, default=str))
            assert row and row[0] not in ("terminal", "privacy_restricted"), (
                row[:2] if row else None
            )
            if row[0] != "accepted":
                return False
            org = flow.execute(
                "SELECT organization_id FROM applications WHERE id=%s", (application,)
            ).fetchone()[0]
            self.owner.execute(
                "SELECT set_config('app.current_tenant_id',%s,false)", (f"org_{org}",)
            )
            return self.owner.execute(
                "SELECT published_generation_id FROM candidate_index_heads WHERE authority_source_id=%s AND published_generation_id IS NOT NULL",
                (row[2],),
            ).fetchone()

        return self.wait(complete)[0]

    def search(self, client, **filters):
        result = self.post(
            client,
            "/api/candidates/semantic-search",
            json={
                "query": "Python backend engineer",
                "use_reranker": False,
                "metadata_filters": filters,
            },
        )
        assert result.status_code == 200, (result.status_code, result.json().get("code"))
        return result.json()["results"]

    def ready(self, port):
        import httpx

        try:
            response = httpx.get(
                f"http://127.0.0.1:{port}/readyz",
                headers={"Authorization": "Bearer fixture-control-token-4d-minimum-32"},
                trust_env=False,
                timeout=2,
            )
            self.file(f"readiness-{port}.json", json.dumps(response.json()))
            return response.status_code == 200
        except httpx.HTTPError:
            return False

    def setup(self):
        from cryptography.hazmat.primitives import serialization
        from cryptography.hazmat.primitives.asymmetric import rsa

        from activekg.candidate_index.contracts import IndexPolicy, canonical_json, sha256
        from activekg.engine.model_config import DEFAULT_GROQ_FAST_MODEL, DEFAULT_GROQ_LARGE_MODEL
        from activekg.extraction.schema import ExtractionResult

        self.bootstrap()
        private = rsa.generate_private_key(public_exponent=65537, key_size=2048)
        pem = private.private_bytes(
            serialization.Encoding.PEM,
            serialization.PrivateFormat.PKCS8,
            serialization.NoEncryption(),
        ).decode()
        public = (
            private.public_key()
            .public_bytes(
                serialization.Encoding.PEM, serialization.PublicFormat.SubjectPublicKeyInfo
            )
            .decode()
        )
        ft = None
        with psycopg.connect(self.fo, autocommit=True) as flow:
            ft = flow.execute("SELECT target_id FROM schema_control.identity").fetchone()[0]
        mt = str(
            self.owner.execute(
                "SELECT target_id FROM activekg_schema_control.target_identity"
            ).fetchone()[0]
        )
        self.start(
            [
                "redis-server",
                "--bind",
                "127.0.0.1",
                "--port",
                str(self.redisport),
                "--save",
                "",
                "--appendonly",
                "no",
            ],
            {},
        )
        revision = "6" * 40
        cache = self.path / "model-cache"
        snapshot = (
            cache / "models--sentence-transformers--all-MiniLM-L6-v2" / "snapshots" / revision
        )
        snapshot.mkdir(parents=True)
        (snapshot / "modules.json").write_text(
            "[]"
        )  # explicit synthetic inference fixture, never real model evidence
        policy = IndexPolicy(
            extraction_schema_sha256=sha256(canonical_json(ExtractionResult.model_json_schema())),
            extraction_prompt_sha256=sha256(
                (self.memory / "activekg/extraction/prompt.py").read_bytes()
            ),
            primary_model_id=DEFAULT_GROQ_FAST_MODEL,
            fallback_model_id=DEFAULT_GROQ_LARGE_MODEL,
            embedding_model_id="all-MiniLM-L6-v2",
            embedding_artifact_revision=revision,
        )
        allowed = json.dumps(
            self.ports
            + [int(conninfo_to_dict(self.fr)["port"]), int(conninfo_to_dict(self.mr)["port"])]
        )
        self.forbidden = self.file("forbidden", "")
        self.model_calls = self.file("model-calls", "")
        self.memory_boot = self.file("memory-process.py", MEMORY_PROCESS)
        models = self.file("models.py", MODEL_BOUNDARY)
        self.me = {
            "ACTIVEKG_DSN": self.mr,
            "ACTIVEKG_SCHEMA_TARGET_ID": mt,
            "ACTIVEKG_SCHEMA_ENVIRONMENT": "development",
            "ACTIVEKG_CONTROL_PLANE_TOKEN": "fixture-control-token-4d-minimum-32",
            "JWT_ENABLED": "true",
            "JWT_ALGORITHM": "RS256",
            "JWT_PUBLIC_KEY": public,
            "JWT_AUDIENCE": "activekg",
            "JWT_ISSUER": "vantahire",
            "LLM_ENABLED": "false",
            "RUN_SCHEDULER": "false",
            "RATE_LIMIT_ENABLED": "false",
            "REDIS_URL": f"redis://127.0.0.1:{self.redisport}/0",
            "CANDIDATE_PRIVACY_FLOW_ISSUER": "vantahire",
            "CANDIDATE_PRIVACY_FLOW_ACTOR_ID": "vantahire-backend",
            "CANDIDATE_PRIVACY_SIGNAL_ISSUER": "signal",
            "CANDIDATE_PRIVACY_SIGNAL_ACTOR_ID": "signal-service",
            "CANDIDATE_INDEX_EMBEDDING_ARTIFACT_REVISION": revision,
            "CANDIDATE_INDEX_POLICY_SHA256": policy.digest,
            "EMBEDDING_BACKEND": "sentence-transformers",
            "EMBEDDING_MODEL": "all-MiniLM-L6-v2",
            "HF_HUB_CACHE": str(cache),
            "FIXTURE_MODEL_BOUNDARY": str(models),
            "FIXTURE_MEMORY_ROOT": str(self.memory),
            "FIXTURE_PORTS": allowed,
            "FIXTURE_API_PORT": str(self.apiport),
            "FIXTURE_FORBIDDEN": str(self.forbidden),
            "FIXTURE_MODEL_CALLS": str(self.model_calls),
        }
        ae = {
            **self.me,
            "CANDIDATE_INDEX_API_ENABLED": "true",
            "CANDIDATE_PRIVACY_INTAKE_ENABLED": "true",
            "ORGANIZATION_CANDIDATE_INTAKE_ENABLED": "true",
            "CANDIDATE_CONSENT_INTAKE_ENABLED": "true",
            "CANDIDATE_PRIVACY_HMAC_ACTIVE_VERSION": "1",
            "CANDIDATE_PRIVACY_HMAC_KEY_V1": base64.b64encode(b"\1" * 32).decode(),
            "ORG_DECISION_INBOX_ENABLED": "true",
            "SOURCED_CANDIDATE_INGEST_MODE": "canonical_only",
        }
        self.run([sys.executable, "scripts/schema_ready.py"], ae, self.memory)
        self.start([sys.executable, "-B", self.memory_boot, "api"], ae, self.memory)
        self.wait(lambda: self.ready(self.apiport))
        for stage, flag, port in [
            ("extraction", "CANDIDATE_INDEX_EXTRACTION_ENABLED", self.extractport),
            ("embedding", "CANDIDATE_INDEX_EMBEDDING_ENABLED", self.embedport),
        ]:
            we = {**self.me, flag: "true", "EXTRACTION_HEALTHCHECK_PORT": str(port)}
            if stage == "extraction":
                we["GROQ_API_KEY"] = "synthetic-fixture-never-a-provider-key"
            self.start([sys.executable, "-B", self.memory_boot, stage], we, self.memory)
            self.wait(lambda port=port: self.ready(port))
        # Storage alias only: the frozen upload/downloader and all byte checking
        # are rebuilt unchanged. The child still uses its strict environment.
        objects = self.path / "objects"
        objects.mkdir(mode=0o700)
        sdk = self.file(
            "storage.mjs",
            """
import fs from 'node:fs';import p from 'node:path';import crypto from 'node:crypto';
const root=ROOT;
class Storage{bucket(bucket){return {file(name){const path=p.join(root,crypto.createHash('sha256').update(bucket+'/'+name).digest('hex'));return {
 save:async bytes=>fs.writeFileSync(path,bytes,{mode:0o600}),download:async()=>[fs.readFileSync(path)],
 delete:async()=>fs.unlinkSync(path),exists:async()=>[fs.existsSync(path)]};}}}}
export {Storage};
""".replace("ROOT", json.dumps(str(objects))),
        )
        self.flow_fence = self.file(
            "flow-fence.mjs",
            """
import net from 'node:net';import dns from 'node:dns';import fs from 'node:fs';import {syncBuiltinESMExports} from 'node:module';
const allowed=new Set(PORTS),f=FILE;
const refuse=()=>{fs.appendFileSync(f,'network\\n');throw Error('fixture_network_refused');};
const original=net.Socket.prototype.connect;
net.Socket.prototype.connect=function(...args){const a=Array.isArray(args[0])?args[0][0]:args[0];
 const host=typeof a==='object'?a.host:args[1],port=typeof a==='object'?a.port:a;
 if(host!=='127.0.0.1'||!allowed.has(Number(port)))return refuse();return original.apply(this,args);};
const lookup=dns.lookup;dns.lookup=(host,...args)=>host==='127.0.0.1'?lookup(host,...args):refuse();syncBuiltinESMExports();
""".replace("PORTS", allowed).replace("FILE", json.dumps(str(self.forbidden))),
        )
        self.run(["npm", "run", "build:server"], {}, timeout=120)
        self.web_build = self.path / "index.mjs"
        (self.path / "node_modules").symlink_to(
            self.flow / "node_modules", target_is_directory=True
        )
        (self.path / "public").symlink_to(self.flow / "dist/public", target_is_directory=True)
        (self.path / "server").symlink_to(self.flow / "dist/server", target_is_directory=True)
        esbuild = self.flow / "node_modules/.bin/esbuild"
        for entry, dest, fmt in [
            ("server/index.ts", self.web_build, "esm"),
            ("server/gcs-storage.ts", self.flow / "dist/candidate-index-gcs.cjs", "cjs"),
        ]:
            self.run(
                [
                    esbuild,
                    entry,
                    "--platform=node",
                    "--packages=external",
                    "--bundle",
                    f"--format={fmt}",
                    f"--outfile={dest}",
                    f"--alias:@google-cloud/storage={sdk}",
                ],
                {},
            )
        self.fe = {
            "DATABASE_URL": self.fr,
            "DATABASE_SSL": "false",
            "PGPOOL_MAX": "5",
            "FLOW_SCHEMA_TARGET_ID": ft,
            "FLOW_SCHEMA_ENVIRONMENT": "development",
            "ACTIVEKG_BASE_URL": self.api_url,
            "VANTAHIRE_JWT_PRIVATE_KEY": pem,
            "VANTAHIRE_JWT_ACTIVE_KID": "v1",
            "FLOW_CANDIDATE_PRIVACY_INTAKE_ENABLED": "true",
            "ORGANIZATION_CANDIDATE_SYNC_ENABLED": "true",
            "CANDIDATE_CONSENT_DELIVERY_ENABLED": "true",
            "ACTIVEKG_SYNC_ENABLED": "true",
            "FLOW_CANDIDATE_INDEX_MODE": "private_primary",
            "SESSION_SECRET": "synthetic-session-secret-4d",
            "PORT": str(self.webport),
            "HOST": "127.0.0.1",
            "SEED_DEFAULTS": "false",
            "EMAIL_AUTOMATION_ENABLED": "false",
            "NOTIFICATION_AUTOMATION_ENABLED": "false",
            "REDIS_URL": f"redis://127.0.0.1:{self.redisport}/0",
            "GCS_PROJECT_ID": "fixture",
            "GCS_BUCKET_NAME": "fixture",
            "GCS_SERVICE_ACCOUNT_KEY": '{"type":"service_account"}',
        }
        self.run(["node", "dist/schema-ready.js"], self.fe)
        self.web = self.start(["node", "--import", self.flow_fence, self.web_build], self.fe)
        import httpx

        def web_ready():
            try:
                return (
                    httpx.get(
                        self.web_url + "/api/csrf-token", trust_env=False, timeout=2
                    ).status_code
                    == 200
                )
            except httpx.HTTPError:
                return False

        self.wait(web_ready)


@pytest.mark.skipif(
    os.getenv("ACTIVEKG_INDEX_XS_PROCESSES") != "1",
    reason="full-process rehearsal separately selected",
)
@pytest.mark.timeout(420)
def test_built_web_real_api_and_workers_publish_private_application(target, tmp_path):  # noqa: F811
    import httpx
    from docx import Document

    stack = ProcessStack(tmp_path, target)
    try:
        stack.setup()
        with psycopg.connect(stack.fo, autocommit=True) as flow:
            suffix = uuid4().hex
            salt = "fixture-salt"
            password = (
                hashlib.scrypt(
                    b"fixture-password", salt=salt.encode(), n=16384, r=8, p=1, dklen=64
                ).hex()
                + "."
                + salt
            )
            recruiter = flow.execute(
                "INSERT INTO users(username,password,role,email_verified,auth_version) VALUES(%s,%s,'recruiter',true,1) RETURNING id",
                (f"lead-{suffix}@fixture.invalid", password),
            ).fetchone()[0]
            org = flow.execute(
                "INSERT INTO organizations(name,slug,is_active) VALUES('Index fixture',%s,true) RETURNING id",
                ("fixture-" + suffix,),
            ).fetchone()[0]
            flow.execute(
                "INSERT INTO organization_members(organization_id,user_id,role,seat_assigned) VALUES(%s,%s,'owner',true)",
                (org, recruiter),
            )
            job = flow.execute(
                "INSERT INTO jobs(organization_id,title,location,type,description,posted_by,is_active,status,slug) VALUES(%s,'Backend engineer','Remote','full-time','Python backend systems',%s,true,'approved',%s) RETURNING id",
                (org, recruiter, "fixture-" + suffix),
            ).fetchone()[0]

            def fresh():
                row = flow.execute(
                    "SELECT status,last_success_at > clock_timestamp()-interval '1 minute' FROM candidate_privacy_sync_state WHERE consumer_name='flow'"
                ).fetchone()
                return row and row[0] == "healthy" and row[1]

            stack.wait(fresh)
            document = Document()
            document.add_paragraph(
                "Backend engineer. Python database development and distributed systems."
            )
            buf = io.BytesIO()
            document.save(buf)
            with httpx.Client(base_url=stack.web_url, trust_env=False, timeout=40) as client:
                token = client.get("/api/csrf-token").json()["token"]
                # Loopback HTTP carries the actual secure double-submit cookie
                # explicitly; the real middleware still checks the same token.
                response = client.post(
                    f"/api/jobs/{job}/apply",
                    data={
                        "name": "Index fixture",
                        "email": f"person-{suffix}@fixture.invalid",
                        "phone": "2025550147",
                        "whatsappConsent": "false",
                    },
                    files={
                        "resume": (
                            "fixture.docx",
                            buf.getvalue(),
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        )
                    },
                    headers={"x-csrf-token": token, "Cookie": "__Host-psifi.x-csrf-token=" + token},
                )
                assert response.status_code in (200, 201), (
                    response.status_code,
                    response.json().get("code"),
                    response.json().get("error"),
                )
                application = response.json()["applicationId"]

                def acknowledged():
                    row = flow.execute(
                        "SELECT d.state,d.error_code FROM candidate_index_outbox o LEFT JOIN candidate_index_delivery_state d USING(outbox_id) WHERE o.application_id=%s",
                        (application,),
                    ).fetchone()
                    stack.file("delivery.json", json.dumps(row))
                    assert row and row[0] not in ("terminal", "privacy_restricted"), row
                    return row[0] == "accepted"

                stack.wait(acknowledged)
                stack.owner.execute(
                    "SELECT set_config('app.current_tenant_id',%s,false)", (f"org_{org}",)
                )

                def published():
                    return (
                        stack.owner.execute(
                            "SELECT count(*) FROM candidate_index_heads WHERE scope_key=%s AND published_generation_id IS NOT NULL",
                            (f"org_{org}",),
                        ).fetchone()[0]
                        == 1
                    )

                stack.wait(published)
                assert stack.model_calls.read_text().splitlines() == ["extract"]
                assert not stack.forbidden.read_bytes()

                def csrf_headers():
                    value = client.get("/api/csrf-token").json()["token"]
                    return {
                        "x-csrf-token": value,
                        "Cookie": "; ".join(f"{k}={v}" for k, v in client.cookies.items()),
                    }

                login = client.post(
                    "/api/login",
                    json={
                        "username": f"lead-{suffix}@fixture.invalid",
                        "password": "fixture-password",
                    },
                    headers=csrf_headers(),
                )
                assert login.status_code == 200
                for filters, expected in [
                    ({}, [application]),
                    ({"source": "other", "org_id": org + 1}, [application]),
                    ({"job_id": job}, [application]),
                    ({"job_id": job + 1}, []),
                ]:
                    found = client.post(
                        "/api/candidates/semantic-search",
                        json={
                            "query": "Python backend engineer",
                            "use_reranker": False,
                            "metadata_filters": filters,
                        },
                        headers=csrf_headers(),
                    )
                    assert found.status_code == 200, (found.status_code, found.json().get("code"))
                    assert [r["applicationId"] for r in found.json()["results"]] == expected
                refused = client.post(
                    "/api/candidates/semantic-search",
                    json={"query": "Python", "metadata_filters": {"country": "test"}},
                    headers=csrf_headers(),
                )
                assert refused.status_code == 422
                assert not stack.forbidden.read_bytes()
    finally:
        stack.close()


@pytest.mark.skipif(
    os.getenv("ACTIVEKG_INDEX_XS_PROCESSES") != "1",
    reason="full-process rehearsal separately selected",
)
@pytest.mark.parametrize("variant", ["anonymous_optout", "saved_optout", "erasure"])
@pytest.mark.timeout(420)
def test_built_web_privacy_request_private_index_and_org_isolation(target, tmp_path, variant):  # noqa: F811
    import httpx

    stack = ProcessStack(tmp_path, target)
    try:
        stack.setup()
        with (
            psycopg.connect(stack.fo, autocommit=True) as flow,
            httpx.Client(base_url=stack.web_url, trust_env=False, timeout=40) as candidate,
            httpx.Client(base_url=stack.web_url, trust_env=False, timeout=40) as anonymous,
            httpx.Client(base_url=stack.web_url, trust_env=False, timeout=40) as lead,
        ):
            org, _, lead_email, job = stack.seed_job(flow)
            _, _, other_email, _ = stack.seed_job(flow)
            user, email = stack.seed_user(flow, "candidate")
            stack.privacy_fresh(flow)
            stack.login(candidate, email)
            assert (
                stack.post(
                    candidate,
                    "/api/candidate/privacy/reauth",
                    json={"password": "fixture-password"},
                ).status_code
                == 200
            )
            request_id = str(uuid4())
            requested = stack.post(
                candidate,
                "/api/candidate/privacy/requests",
                json={
                    "requestId": request_id,
                    "action": "request_erasure"
                    if variant == "erasure"
                    else "withdraw_global_matching",
                },
            )
            assert requested.status_code == 202

            def delivered():
                states = candidate.get("/api/candidate/privacy/status").json()["requests"]
                row = next((r for r in states if r["requestId"] == request_id), None)
                return (
                    row
                    and row["deliveryStatus"] == "delivered"
                    and row["decision"] == ("block_all" if variant == "erasure" else "block_global")
                )

            stack.wait(delivered)
            # Authority was created by the real Flow route and web processor;
            # no directive, receipt or privacy decision is planted in Memory.
            before = flow.execute("SELECT count(*) FROM applications").fetchone()[0]
            resume_id = None
            if variant == "saved_optout":
                # An existing saved DOCX with no extracted text exercises the
                # original-byte sender, actual isolated DOCX parser and hash pin.
                name = "resumes/" + uuid4().hex + ".docx"
                path = (
                    stack.path
                    / "objects"
                    / hashlib.sha256(("fixture/" + name).encode()).hexdigest()
                )
                path.write_bytes(stack.resume()[1])
                path.chmod(0o600)
                resume_id = flow.execute(
                    "INSERT INTO candidate_resumes(user_id,label,gcs_path,extracted_text,updated_at) VALUES(%s,'fixture.docx',%s,NULL,date_trunc('milliseconds',clock_timestamp())) RETURNING id",
                    (user, "gs://fixture/" + name),
                ).fetchone()[0]
            response = stack.apply(
                candidate if variant == "saved_optout" else anonymous,
                job,
                email,
                resume_id=resume_id,
            )
            stack.file(
                "apply-result.json",
                json.dumps({"status": response.status_code, "code": response.json().get("code")}),
            )
            if variant == "erasure":
                # Frozen Flow public-apply maps privacy denials to 503 plus a
                # fixed code; 451 belongs to Memory and consent endpoints.
                assert response.status_code == 503
                assert response.json()["code"] == "candidate_privacy_restricted"
                assert flow.execute("SELECT count(*) FROM applications").fetchone()[0] == before
                assert not stack.model_calls.read_bytes()
            else:
                assert response.status_code == 201, (
                    response.status_code,
                    response.json().get("code"),
                )
                application = response.json()["applicationId"]
                stack.published(flow, application)
                assert (
                    flow.execute(
                        "SELECT count(*) FROM application_graph_sync_jobs WHERE application_id=%s",
                        (application,),
                    ).fetchone()[0]
                    == 0
                )
                assert flow.execute(
                    "SELECT user_id FROM applications WHERE id=%s", (application,)
                ).fetchone()[0] == (user if variant == "saved_optout" else None)
                assert (
                    stack.owner.execute(
                        "SELECT count(*) FROM candidate_index_read_public(100)"
                    ).fetchone()[0]
                    == 0
                )
                assert (
                    stack.owner.execute("SELECT count(*) FROM global_candidates").fetchone()[0] == 0
                )
                stack.login(lead, lead_email)
                assert [r["applicationId"] for r in stack.search(lead, job_id=job)] == [application]
                with httpx.Client(base_url=stack.web_url, trust_env=False, timeout=40) as other:
                    stack.login(other, other_email)
                    assert stack.search(other, org_id=org, job_id=job) == []
                # The retained global writer is still refused for this identity.
                _, _, _, second_job = stack.seed_job(
                    flow,
                    org=org,
                    recruiter=flow.execute(
                        "SELECT posted_by FROM jobs WHERE id=%s", (job,)
                    ).fetchone()[0],
                )
                denied = stack.apply(lead, second_job, email, recruiter=True)
                assert denied.status_code == 503
                assert denied.json()["code"] == "candidate_privacy_restricted"
                assert stack.model_calls.read_text().splitlines() == ["extract"]
            assert not stack.forbidden.read_bytes()
    finally:
        stack.close()


@pytest.mark.skipif(
    os.getenv("ACTIVEKG_INDEX_XS_PROCESSES") != "1",
    reason="full-process rehearsal separately selected",
)
@pytest.mark.timeout(420)
def test_built_web_cutover_preserves_all_three_nonadopter_handlers(target, tmp_path):  # noqa: F811
    import httpx

    stack = ProcessStack(tmp_path, target)
    try:
        stack.setup()
        stack.restart_web(FLOW_CANDIDATE_INDEX_MODE="dual", BULK_RESUME_IMPORT_ENABLED="true")
        with (
            psycopg.connect(stack.fo, autocommit=True) as flow,
            httpx.Client(base_url=stack.web_url, trust_env=False, timeout=40) as anonymous,
            httpx.Client(base_url=stack.web_url, trust_env=False, timeout=40) as lead,
        ):
            org, recruiter, email, job = stack.seed_job(flow)
            stack.privacy_fresh(flow)

            def queue(application):
                return flow.execute(
                    "SELECT status,attempts FROM application_graph_sync_jobs WHERE application_id=%s",
                    (application,),
                ).fetchone()

            dual = stack.apply(anonymous, job, f"dual-{uuid4().hex}@fixture.invalid")
            assert dual.status_code == 201
            dual_id = dual.json()["applicationId"]
            assert queue(dual_id) is not None  # real public-apply legacy bridge, not its predicate
            stack.published(flow, dual_id)
            stack.login(lead, email)
            assert [r["applicationId"] for r in stack.search(lead, job_id=job)].count(dual_id) == 1
            stack.restart_web(FLOW_CANDIDATE_INDEX_MODE="private_primary")
            stack.privacy_fresh(flow)
            primary = stack.apply(anonymous, job, f"primary-{uuid4().hex}@fixture.invalid")
            assert primary.status_code == 201
            primary_id = primary.json()["applicationId"]
            assert queue(primary_id) is None
            stack.published(flow, primary_id)

            # A stale queued row for the exact adopted application reaches the
            # real running web timer. It must settle managed without any legacy
            # parent/chunk write. Both overlapping lock orders are additionally
            # exercised by the real-PostgreSQL worker/catchup race matrix.
            flow.execute(
                "INSERT INTO application_graph_sync_jobs(application_id,organization_id,job_id,effective_recruiter_id,activekg_tenant_id) VALUES(%s,%s,%s,%s,%s)",
                (primary_id, org, job, recruiter, f"org_{org}"),
            )
            stack.wait(lambda: queue(primary_id)[0] == "dead_letter")
            assert (
                flow.execute(
                    "SELECT last_error FROM application_graph_sync_jobs WHERE application_id=%s",
                    (primary_id,),
                ).fetchone()[0]
                == "candidate_index_managed"
            )
            assert (
                stack.owner.execute(
                    "SELECT count(*) FROM nodes WHERE tenant_id=%s AND metadata->>'application_id'=%s",
                    (f"org_{org}", str(primary_id)),
                ).fetchone()[0]
                == 0
            )

            # The three retained writers execute through the built handlers with
            # real organization/seat/CSRF and global-use admission. They never
            # inherit 4D ownership from the source application or resume path.
            added = stack.apply(
                lead, job, f"recruiter-{uuid4().hex}@fixture.invalid", recruiter=True
            )
            assert added.status_code == 201
            add_id = added.json()["applicationId"]
            assert queue(add_id) is not None
            _, _, _, target_job = stack.seed_job(flow, org=org, recruiter=recruiter)
            cloned = stack.post(
                lead,
                "/api/candidates/move-to-job",
                json={"sourceApplicationId": primary_id, "targetJobId": target_job},
            )
            assert cloned.status_code == 201, (cloned.status_code, cloned.json().get("error"))
            clone_id = cloned.json()["applicationId"]
            assert clone_id != primary_id and queue(clone_id) is not None

            batch = flow.execute(
                "INSERT INTO resume_import_batches(organization_id,job_id,uploaded_by_user_id,status,file_count,processed_count,ready_count) VALUES(%s,%s,%s,'ready_for_review',1,1,1) RETURNING id",
                (org, job, recruiter),
            ).fetchone()[0]
            locator = flow.execute(
                "SELECT resume_url FROM applications WHERE id=%s", (add_id,)
            ).fetchone()[0]
            item = flow.execute(
                "INSERT INTO resume_import_items(batch_id,organization_id,job_id,uploaded_by_user_id,original_filename,gcs_path,extracted_text,extraction_method,parsed_name,parsed_email,parsed_phone,status) VALUES(%s,%s,%s,%s,'fixture.docx',%s,%s,'native_text','Import fixture',%s,'2025550147','processed') RETURNING id",
                (
                    batch,
                    org,
                    job,
                    recruiter,
                    locator,
                    stack.resume()[0],
                    f"bulk-{uuid4().hex}@fixture.invalid",
                ),
            ).fetchone()[0]
            finalized = stack.post(
                lead,
                f"/api/jobs/{job}/bulk-resume-import/{batch}/finalize",
                json={"itemIds": [item]},
            )
            assert finalized.status_code == 200
            assert len(finalized.json()["finalized"]) == 1, finalized.json().get("needsReview")
            bulk_id = finalized.json()["finalized"][0]["applicationId"]
            assert queue(bulk_id) is not None
            for app_id in (add_id, clone_id, bulk_id):
                assert (
                    flow.execute(
                        "SELECT flow_candidate_index_managed_application(%s,%s)", (org, app_id)
                    ).fetchone()[0]
                    is False
                )
                assert (
                    flow.execute(
                        "SELECT count(*) FROM candidate_index_outbox WHERE application_id=%s",
                        (app_id,),
                    ).fetchone()[0]
                    == 0
                )

            # Unchanged flag-off behavior: the eligible handler still commits
            # the application but must not enqueue either indexing lane.
            stack.restart_web(ACTIVEKG_SYNC_ENABLED="false")
            stack.privacy_fresh(flow)
            disabled = stack.apply(
                lead, job, f"disabled-{uuid4().hex}@fixture.invalid", recruiter=True
            )
            assert disabled.status_code == 201
            assert queue(disabled.json()["applicationId"]) is None
            assert not stack.forbidden.read_bytes()
    finally:
        stack.close()


@pytest.mark.skipif(
    os.getenv("ACTIVEKG_INDEX_XS_PROCESSES") != "1",
    reason="full-process rehearsal separately selected",
)
@pytest.mark.timeout(420)
def test_built_web_consent_publication_replacement_and_withdrawal_keep_private_search(
    target,  # noqa: F811
    tmp_path,
):
    import httpx

    stack = ProcessStack(tmp_path, target)
    try:
        stack.setup()
        with (
            psycopg.connect(stack.fo, autocommit=True) as flow,
            httpx.Client(base_url=stack.web_url, trust_env=False, timeout=40) as candidate,
            httpx.Client(base_url=stack.web_url, trust_env=False, timeout=40) as lead,
        ):
            org, _, lead_email, job = stack.seed_job(flow)
            user, email = stack.seed_user(flow, "candidate")
            stack.privacy_fresh(flow)
            stack.login(candidate, email)
            applied = stack.apply(candidate, job, email)
            assert applied.status_code == 201
            application = applied.json()["applicationId"]
            stack.published(flow, application)
            resume_id = str(
                flow.execute(
                    "SELECT resume_version_id FROM application_resume_versions WHERE application_id=%s",
                    (application,),
                ).fetchone()[0]
            )
            assert (
                stack.post(
                    candidate,
                    "/api/candidate/privacy/reauth",
                    json={"password": "fixture-password"},
                ).status_code
                == 200
            )
            contract = candidate.get("/api/candidate/consent").json()

            def public_rows():
                return [
                    row[0]
                    for row in stack.owner.execute(
                        "SELECT * FROM candidate_index_read_public(100)"
                    ).fetchall()
                ]

            def settled(version, action):
                status = candidate.get("/api/candidate/consent").json()
                effective = status.get("effective") or {}
                return effective.get("version") == version and effective.get("action") == action

            with psycopg.connect(stack.mr, autocommit=True) as runtime:
                with pytest.raises(psycopg.errors.InsufficientPrivilege):
                    runtime.execute("SELECT * FROM candidate_index_read_public(100)")

            first_generation = None
            for version, resume in [(1, None), (2, resume_id)]:
                body = {
                    "request_id": str(uuid4()),
                    "expected_version": version - 1,
                    "purpose": contract["purpose"],
                    "copy_version": contract["copy_version"],
                    "copy_sha256": contract["copy_sha256"],
                    "profile": {
                        "display_name": "Approved fixture",
                        "headline": f"Python backend version {version}",
                        "location": "Remote",
                        "skills": ["Python"],
                        "linkedin": None,
                    },
                    "resume_version_id": resume,
                }
                granted = stack.post(candidate, "/api/candidate/consent/grant", json=body)
                assert granted.status_code in (200, 202), (
                    granted.status_code,
                    granted.json().get("code"),
                )
                stack.wait(lambda version=version: settled(version, "grant"))
                # Current authority changes at the receiver commit; an older
                # approved generation cannot remain visible while its replacement
                # is being built. A completed new generation may already exist.
                assert all(r["generation_id"] != first_generation for r in public_rows())
                replay = stack.post(candidate, "/api/candidate/consent/grant", json=body)
                assert replay.status_code in (200, 202) and replay.json()["replayed"] is True
                rows = stack.wait(lambda: public_rows())
                assert len(rows) == 1 and rows[0]["approved_profile"] == body["profile"]
                first_generation = rows[0]["generation_id"]
                # No public search route or legacy ready profile is published.
                assert (
                    stack.owner.execute(
                        "SELECT count(*) FROM global_candidates WHERE embedding IS NOT NULL"
                    ).fetchone()[0]
                    == 0
                )
                private_head = stack.published(flow, application)
                assert private_head is not None
            withdrawn = stack.post(
                candidate,
                "/api/candidate/consent/withdraw",
                json={"request_id": str(uuid4()), "expected_version": 2},
            )
            assert withdrawn.status_code in (200, 202)
            stack.wait(lambda: settled(3, "withdraw"))
            assert public_rows() == []
            stack.login(lead, lead_email)
            assert [r["applicationId"] for r in stack.search(lead, job_id=job)] == [application]
            assert (
                flow.execute(
                    "SELECT count(*) FROM candidate_consent_events WHERE user_id=%s", (user,)
                ).fetchone()[0]
                == 3
            )
            assert not stack.forbidden.read_bytes()
    finally:
        stack.close()
