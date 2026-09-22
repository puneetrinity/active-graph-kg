"""One-attempt HTTP/child transports and generation orchestration, no live models."""

import asyncio
import base64
import io
import json
import os
import subprocess
import sys
import threading
import time
from dataclasses import asdict, replace
from pathlib import Path
from types import SimpleNamespace

import httpx
import pytest

from activekg.candidate_index import processing, providers
from activekg.candidate_index.contracts import (
    IndexContractError,
    IndexPolicy,
    IndexTunables,
    canonical_json,
    chunk_manifest,
    sha256,
)
from activekg.candidate_index.processing import (
    GenerationRuntime,
    GenerationWorker,
    profile_extraction,
    resume_extraction,
    worker_configuration,
)
from activekg.candidate_index.providers import (
    WorkRefused,
    extract_once,
    parse_original,
    retry_delay,
)
from tests.test_candidate_index import POLICY

TEXT = "Backend engineer. Python and SQL. Distributed systems."


@pytest.mark.parametrize("case", ["clean", "network", "overflow", "failure", "insecure"])
def test_ml_import_is_bounded_before_stdin_and_processing_is_zero(tmp_path, case):
    scratch = tmp_path / "import"
    scratch.mkdir(mode=0o700)
    if case == "insecure":
        scratch.chmod(0o755)
    script = r"""
import builtins, importlib.util, io, json, os, resource, signal, socket, sys, types
spec = importlib.util.spec_from_file_location('bounded_child', sys.argv[1])
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)
case = sys.argv[2]
original = builtins.__import__
imported = False
class Input:
    @property
    def buffer(self): return self
    def read(self, size):
        assert imported
        assert resource.getrlimit(resource.RLIMIT_FSIZE) == (0, 0)
        return b'{}'  # malformed payload must still be read only after the fence
sys.stdin = Input()
def load(name, *args, **kwargs):
    global imported
    if name in ('torch', 'sentence_transformers'):
        assert resource.getrlimit(resource.RLIMIT_FSIZE) == (1048576, 1048576)
        assert os.stat(os.getcwd()).st_mode & 0o777 == 0o700
        if name == 'sentence_transformers':
            if case == 'network':
                try: socket.getaddrinfo('synthetic.invalid', 443)
                except module.WorkRefused: pass
            if case == 'failure': raise ImportError('synthetic')
            signal.signal(signal.SIGXFSZ, signal.SIG_IGN)
            with open('trusted-template', 'wb', buffering=0) as output:
                if case == 'overflow':
                    output.write(b'x' * 1048576)
                    try: output.write(b'x')
                    except OSError as exc:
                        assert exc.errno == 27
                        raise ImportError('bounded') from None
                    raise AssertionError('limit not enforced')
                output.write(b'trusted')
            imported = True
        return types.ModuleType(name)
    return original(name, *args, **kwargs)
builtins.__import__ = load
module._run_child('embed')
"""
    result = subprocess.run(
        [sys.executable, "-I", "-B", "-c", script, providers.__file__, case],
        env={**providers._child_environment(), "CANDIDATE_INDEX_IMPORT_DIR": str(scratch)},
        capture_output=True,
        timeout=5,
    )
    assert result.returncode == 0 and not result.stderr
    assert json.loads(result.stdout) == {
        "error": "incomplete" if case == "clean" else "provider_unavailable"
    }


@pytest.mark.parametrize("outcome", ["success", "failure", "cancel"])
def test_parent_removes_ml_scratch_after_reaped_child(monkeypatch, outcome):
    observed = []

    async def child(mode, payload, *, deadline, scratch):
        path = Path(scratch)
        observed.append(path)
        assert path.stat().st_mode & 0o777 == 0o700
        (path / "trusted-template").write_bytes(b"synthetic")
        if outcome == "failure":
            raise WorkRefused("provider_unavailable")
        if outcome == "cancel":
            raise asyncio.CancelledError()
        return []

    monkeypatch.setattr(providers, "_child_work", child)
    try:
        asyncio.run(providers.child_work("embed", {}, deadline=time.monotonic() + 1))
    except (WorkRefused, asyncio.CancelledError):
        assert outcome != "success"
    assert len(observed) == 1 and not observed[0].exists()


RESULT = {"current_title": "Backend engineer", "skills_raw": ["Python", "SQL"], "confidence": 0.9}


def wire(result=RESULT, *, model=POLICY.primary_model_id, finish="stop"):
    return {
        "model": model,
        "choices": [
            {
                "finish_reason": finish,
                "message": {"role": "assistant", "content": json.dumps(result)},
            }
        ],
    }


class Stream(httpx.AsyncByteStream):
    def __init__(self, blocks, *, stall=False):
        self.blocks = blocks
        self.stall = stall
        self.closed = False

    async def __aiter__(self):
        for block in self.blocks:
            yield block
        if self.stall:
            await asyncio.Event().wait()

    async def aclose(self):
        self.closed = True


def response(data, *, headers=None, status=200):
    return httpx.Response(
        status,
        headers=headers or {"content-type": "application/json"},
        stream=Stream([json.dumps(data).encode()]),
    )


async def admitted():
    return None


def invoke(handler, *, before=admitted, deadline=None, text=TEXT):
    return asyncio.run(
        extract_once(
            text,
            POLICY.primary_model_id,
            api_key="synthetic-loopback-key",
            deadline=deadline or time.monotonic() + 2,
            before_send=before,
            transport=httpx.MockTransport(handler),
        )
    )


def test_single_request_frozen_prompt_and_privacy_immediately_before_transport():
    events = []

    async def proof():
        events.append("privacy")

    def handler(request):
        events.append("http")
        assert request.url == providers.GROQ_ENDPOINT
        assert request.headers["authorization"] == "Bearer synthetic-loopback-key"
        payload = json.loads(request.content)
        from activekg.extraction.prompt import build_extraction_prompt

        system, prompt = build_extraction_prompt(TEXT)
        assert payload == {
            "model": POLICY.primary_model_id,
            "messages": [
                {"role": "system", "content": system},
                {"role": "user", "content": prompt},
            ],
            "temperature": 0.1,
            "max_tokens": 1024,
            "stop": [],
            "stream": False,
        }
        return response(wire())

    assert invoke(handler, before=proof) == RESULT
    assert events == ["privacy", "http"]


def test_privacy_refusal_and_prompt_truncation_make_zero_requests():
    calls = []

    def handler(request):
        calls.append(request)
        return response(wire())

    async def denied():
        raise WorkRefused("privacy_restricted")

    with pytest.raises(WorkRefused, match="privacy_restricted"):
        invoke(handler, before=denied)
    from activekg.extraction.prompt import EXTRACTION_MAX_INPUT_CHARS

    with pytest.raises(WorkRefused, match="incomplete"):
        invoke(handler, text="x" * (EXTRACTION_MAX_INPUT_CHARS + 1))
    assert calls == []


@pytest.mark.parametrize("status", [301, 302, 400, 401, 403, 500, 503])
def test_http_errors_are_closed_single_attempts(status, caplog):
    calls = []

    def handler(request):
        calls.append(request)
        return response({"error": "SOURCE-SENTINEL"}, status=status)

    with pytest.raises(WorkRefused, match="provider_unavailable") as error:
        invoke(handler)
    assert len(calls) == 1 and "SOURCE-SENTINEL" not in str(error.value) + caplog.text


@pytest.mark.parametrize(
    "headers,expected",
    [
        ({"retry-after": "2"}, 2000),
        ({"x-ratelimit-remaining-tokens": "0", "x-ratelimit-reset-tokens": "7.66s"}, 7660),
        (
            {
                "retry-after": "2",
                "x-ratelimit-remaining-requests": "0",
                "x-ratelimit-reset-requests": "1m1.2s",
            },
            61200,
        ),
        ({"x-ratelimit-remaining-tokens": "3", "x-ratelimit-reset-tokens": "50s"}, 1000),
        ({"retry-after": "nan"}, 1000),
        ({"retry-after": "nonsense"}, 1000),
        ({"retry-after": "Thu, 01 Jan 1970 00:00:07 GMT"}, 7000),
    ],
)
def test_documented_retry_and_reset_units(headers, expected):
    assert retry_delay(headers, wall_time=0) == expected


def test_retry_longer_than_deadline_is_not_clipped_and_no_inline_retry():
    calls = []

    def handler(request):
        calls.append(request)
        return response({}, headers={"retry-after": "7"}, status=429)

    with pytest.raises(WorkRefused, match="dispatch_exhausted"):
        invoke(handler)
    assert len(calls) == 1
    with pytest.raises(WorkRefused, match="dispatch_exhausted"):
        retry_delay({"retry-after": "121"})
    with pytest.raises(WorkRefused, match="rate_limited") as error:
        invoke(handler, deadline=time.monotonic() + 20)
    assert error.value.retry_ms == 7000 and len(calls) == 2


@pytest.mark.parametrize(
    "data",
    [
        {},
        [],
        wire(model="other-model"),
        wire(finish="length"),
        wire(finish="tool_calls"),
        {"model": POLICY.primary_model_id, "choices": []},
        {"model": POLICY.primary_model_id, "choices": [{"finish_reason": "stop", "message": None}]},
    ],
)
def test_wrong_or_partial_envelope_refuses(data):
    with pytest.raises(WorkRefused, match="incomplete"):
        invoke(lambda _: response(data))


@pytest.mark.parametrize("text", ['{"confidence":NaN}', '{"x":1,"x":2}', "```json\n{}\n```", "[]"])
def test_no_json_repair_nonfinite_or_duplicate_keys(text):
    value = wire()
    value["choices"][0]["message"]["content"] = text
    with pytest.raises(WorkRefused, match="incomplete"):
        invoke(lambda _: response(value))


def test_stream_overflow_and_stall_are_cancelled_not_detached():
    oversized = Stream([b"x" * 40000, b"x" * 40000])
    with pytest.raises(WorkRefused, match="incomplete"):
        invoke(
            lambda _: httpx.Response(
                200, headers={"content-type": "application/json"}, stream=oversized
            )
        )
    assert oversized.closed
    stalled = Stream([b"{"], stall=True)
    with pytest.raises(WorkRefused, match="provider_timeout"):
        invoke(
            lambda _: httpx.Response(
                200, headers={"content-type": "application/json"}, stream=stalled
            ),
            deadline=time.monotonic() + 0.1,
        )
    assert stalled.closed


def test_task_cancellation_closes_actual_response_stream():
    async def run():
        stalled = Stream([b"{"], stall=True)
        transport = httpx.MockTransport(
            lambda _: httpx.Response(
                200, headers={"content-type": "application/json"}, stream=stalled
            )
        )
        task = asyncio.create_task(
            extract_once(
                TEXT,
                POLICY.primary_model_id,
                api_key="synthetic",
                deadline=time.monotonic() + 5,
                before_send=admitted,
                transport=transport,
            )
        )
        await asyncio.sleep(0.02)
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
        assert stalled.closed

    asyncio.run(run())


@pytest.mark.parametrize(
    "delta,code",
    [
        ({"confidence": 0.64}, "low_confidence"),
        ({"confidence": float("inf")}, "incomplete"),
        ({"confidence": True}, "incomplete"),
        ({"confidence": "0.9"}, "incomplete"),
        ({"skills_normalized": [123]}, "incomplete"),
        ({"skills_raw": ["x"] * 100}, "incomplete"),
        ({"current_title": "x\x00y"}, "incomplete"),
        ({"name": "hidden-field"}, "incomplete"),
        ({"location": {"secret": "hidden-field"}}, "incomplete"),
        ({"years_by_skill": {"Python": float("nan")}}, "incomplete"),
        ({"total_years_experience": -3}, "incomplete"),
        ({"total_years_experience": 10**1000}, "incomplete"),
        ({"total_years_experience": "7-3"}, "incomplete"),
        ({"current_title": False}, "incomplete"),
        ({"seniority": "invented"}, "incomplete"),
    ],
)
def test_both_model_tiers_have_identical_strict_quality_gate(delta, code):
    for model in (POLICY.primary_model_id, POLICY.fallback_model_id):
        with pytest.raises(WorkRefused, match=code):
            resume_extraction({**RESULT, **delta}, TEXT, POLICY, model)


def test_empty_skills_valid_with_professional_evidence_but_name_only_is_not():
    good = resume_extraction(
        {"current_title": "Backend engineer", "confidence": 0.8},
        TEXT,
        POLICY,
        POLICY.primary_model_id,
    )
    assert good.namespace["skills_normalized"] == []
    assert good.namespace["location"] is None  # Whole derived namespace clears missing fields.
    assert good.evidence == [{"field": "current_title", "text": "Backend engineer"}]
    for raw in ({"confidence": 1.0}, {"current_title": "Unrelated title", "confidence": 0.9}):
        with pytest.raises(WorkRefused, match="incomplete"):
            resume_extraction(raw, TEXT, POLICY, POLICY.primary_model_id)


def test_profile_only_is_deterministic_without_name_location_or_url_chunks():
    profile = {
        "display_name": "Synthetic Person",
        "headline": "Backend engineer",
        "skills": ["Python"],
        "location": "Synthetic City",
        "linkedin": None,
    }
    result = profile_extraction(profile, POLICY)
    assert result.model_id == "deterministic-profile:v1" and result.confidence == 1
    assert result.namespace == {"headline": "Backend engineer", "skills": ["Python"]}
    assert "Synthetic" not in result.professional_text
    assert profile["display_name"] == "Synthetic Person"
    with pytest.raises(WorkRefused, match="incomplete"):
        profile_extraction({**profile, "headline": "", "skills": []}, POLICY)
    with pytest.raises(WorkRefused, match="chunk_overflow"):
        profile_extraction(profile, replace(POLICY, max_chunks=1, chunk_characters=10))


def docx_bytes():
    from docx import Document

    document = Document()
    document.add_paragraph("Backend engineer")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Python SQL"
    document.sections[0].header.paragraphs[0].text = "Distributed systems"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def pdf_bytes():
    # Minimal real one-page PDF; correct offsets and xref, no extra dependency.
    stream = b"BT /F1 12 Tf 30 100 Td (Backend engineer Python) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    raw = bytearray(b"%PDF-1.4\n")
    offsets = []
    for index, obj in enumerate(objects, 1):
        offsets.append(len(raw))
        raw.extend(str(index).encode() + b" 0 obj\n" + obj + b"\nendobj\n")
    start = len(raw)
    raw.extend(b"xref\n0 6\n0000000000 65535 f \n")
    for offset in offsets:
        raw.extend(f"{offset:010} 00000 n \n".encode())
    raw.extend(
        b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n" + str(start).encode() + b"\n%%EOF\n"
    )
    return bytes(raw)


@pytest.mark.parametrize(
    "factory,expected",
    [
        (docx_bytes, ["Backend engineer", "Python SQL", "Distributed systems"]),
        (pdf_bytes, ["Backend engineer Python"]),
    ],
)
def test_real_parser_child_complete_pdf_and_docx_tables_headers(factory, expected):
    text = asyncio.run(
        parse_original(base64.b64encode(factory()).decode(), deadline=time.monotonic() + 10)
    )
    assert all(value in text for value in expected)


@pytest.mark.parametrize(
    "raw,code",
    [
        (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1binary", "unsupported_format"),
        (b"%PDF-1.4\ntruncated", "incomplete"),
        (b"PK\x03\x04not-zip", "incomplete"),
        (b"html", "unsupported_format"),
    ],
)
def test_original_binary_doc_truncation_and_bad_formats_refuse(raw, code):
    with pytest.raises(WorkRefused, match=code):
        asyncio.run(parse_original(base64.b64encode(raw).decode(), deadline=time.monotonic() + 10))


@pytest.mark.parametrize("mode", ["parse", "embed", "rerank"])
def test_child_receives_no_credentials_and_is_killed_reaped_on_timeout(monkeypatch, mode):
    real = asyncio.create_subprocess_exec
    children = []
    captured = []

    async def launch(*args, **kwargs):
        captured.append((args, kwargs))
        # Substitute only the child process boundary with a stuck parser.
        child = await real(args[0], "-I", "-c", "import time; time.sleep(20)", **kwargs)
        children.append(child)
        return child

    monkeypatch.setenv("GROQ_API_KEY", "DO-NOT-INHERIT")
    monkeypatch.setenv("ACTIVEKG_DSN", "DO-NOT-INHERIT")
    monkeypatch.setattr(asyncio, "create_subprocess_exec", launch)
    with pytest.raises(WorkRefused, match="provider_timeout"):
        asyncio.run(
            providers.child_work(mode, {"original_bytes": "YQ=="}, deadline=time.monotonic() + 0.1)
        )
    assert children[0].returncode is not None
    assert "DO-NOT-INHERIT" not in repr(captured)
    assert "-I" in captured[0][0] and captured[0][1]["env"]["HF_HUB_OFFLINE"] == "1"
    with pytest.raises(ProcessLookupError):
        os.kill(children[0].pid, 0)


def test_exact_cached_revision_missing_is_zero_download(tmp_path):
    with pytest.raises(WorkRefused, match="provider_unavailable"):
        providers.cached_embedding_path("all-MiniLM-L6-v2", "a" * 40, cache_dir=tmp_path)
    path = tmp_path / "models--sentence-transformers--all-MiniLM-L6-v2" / "snapshots" / ("a" * 40)
    path.mkdir(parents=True)
    (path / "modules.json").write_text("[]")
    assert providers.cached_embedding_path("all-MiniLM-L6-v2", "a" * 40, cache_dir=tmp_path) == path
    with pytest.raises(WorkRefused, match="policy_mismatch"):
        providers.cached_embedding_path("all-MiniLM-L6-v2", "main", cache_dir=tmp_path)


def test_query_encoder_uses_pinned_space_with_2000_character_bound(monkeypatch):
    calls = []
    monkeypatch.setattr(
        providers, "cached_embedding_path", lambda m, r: calls.append((m, r)) or Path("/fixed")
    )

    async def child(mode, payload, **kwargs):
        calls.append((mode, payload))
        return [[1.0] + [0.0] * 383]

    monkeypatch.setattr(providers, "child_work", child)
    result = asyncio.run(
        providers.embed_query_once(
            "x" * 2000,
            POLICY.embedding_model_id,
            POLICY.embedding_artifact_revision,
            deadline=time.monotonic() + 2,
        )
    )
    assert result == [1.0] + [0.0] * 383
    assert calls[0] == (POLICY.embedding_model_id, POLICY.embedding_artifact_revision)
    assert calls[1] == (
        "embed",
        {"texts": ["x" * 2000], "path": "/fixed", "revision": POLICY.embedding_artifact_revision},
    )
    with pytest.raises(WorkRefused, match="incomplete"):
        asyncio.run(
            providers.embed_query_once("x" * 2001, "model", "a" * 40, deadline=time.monotonic() + 2)
        )
    assert len(calls) == 2


@pytest.mark.parametrize("scores", [[1.0, -2.0], [1.0], [True, 1.0], [float("nan"), 1.0]])
def test_reranker_fixed_offline_cache_and_finite_exact_cardinality(tmp_path, monkeypatch, scores):
    import huggingface_hub.constants

    monkeypatch.setattr(huggingface_hub.constants, "HF_HUB_CACHE", str(tmp_path))
    root = tmp_path / "models--cross-encoder--ms-marco-MiniLM-L-6-v2"
    (root / "refs").mkdir(parents=True)
    (root / "refs" / "main").write_text("a" * 40)
    path = root / "snapshots" / ("a" * 40)
    path.mkdir(parents=True)
    (path / "config.json").write_text("{}")

    async def child(mode, payload, **kwargs):
        assert mode == "rerank" and payload == {
            "path": str(path),
            "revision": "a" * 40,
            "query": "Python",
            "documents": ["one", "two"],
        }
        return scores

    monkeypatch.setattr(providers, "child_work", child)

    async def run():
        return await providers.rerank_once("Python", ["one", "two"], deadline=time.monotonic() + 2)

    if scores == [1.0, -2.0]:
        assert asyncio.run(run()) == scores
    else:
        with pytest.raises(WorkRefused, match="incomplete"):
            asyncio.run(run())


def test_reranker_missing_cache_is_fallback_not_network(monkeypatch, tmp_path):
    import huggingface_hub.constants

    monkeypatch.setattr(huggingface_hub.constants, "HF_HUB_CACHE", str(tmp_path))

    async def forbidden(*args, **kwargs):
        pytest.fail("No child and no download when the local snapshot is absent")

    monkeypatch.setattr(providers, "child_work", forbidden)
    with pytest.raises(WorkRefused, match="provider_unavailable"):
        asyncio.run(providers.rerank_once("Python", ["one"], deadline=time.monotonic() + 2))


@pytest.mark.parametrize("network_probe", [False, True])
def test_reranker_child_uses_exact_frozen_model_options_with_network_fenced(
    tmp_path, network_probe
):
    path = tmp_path / "models--cross-encoder--ms-marco-MiniLM-L-6-v2" / "snapshots" / ("a" * 40)
    path.mkdir(parents=True)
    scratch = tmp_path / "import"
    scratch.mkdir(mode=0o700)
    script = """
import importlib.util, sys, types, socket
spec = importlib.util.spec_from_file_location('bounded_child', sys.argv[1])
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)
torch = types.ModuleType('torch')
torch.float32 = 'fp32'
sys.modules['torch'] = torch
class CrossEncoder:
    def __init__(self, path, **kwargs):
        assert kwargs == {'device': 'cpu', 'automodel_args': {'device_map': None, 'low_cpu_mem_usage': False, 'dtype': 'fp32'}}
        if sys.argv[2] == 'probe':
            try: socket.getaddrinfo('not-a-real-destination.invalid', 443)
            except module.WorkRefused: pass
    def predict(self, pairs):
        assert pairs == [['Python', 'one'], ['Python', 'two']]
        return types.SimpleNamespace(tolist=lambda: [1.0, -2.0])
stub = types.ModuleType('sentence_transformers')
stub.CrossEncoder = CrossEncoder
sys.modules[stub.__name__] = stub
module._run_child('rerank')
"""
    payload = {
        "path": str(path),
        "revision": "a" * 40,
        "query": "Python",
        "documents": ["one", "two"],
    }
    result = subprocess.run(
        [
            sys.executable,
            "-I",
            "-B",
            "-c",
            script,
            providers.__file__,
            "probe" if network_probe else "clean",
        ],
        input=json.dumps(payload).encode(),
        capture_output=True,
        timeout=5,
        env={**providers._child_environment(), "CANDIDATE_INDEX_IMPORT_DIR": str(scratch)},
    )
    assert result.returncode == 0 and not result.stderr
    output = json.loads(result.stdout)
    output.pop("fence", None)
    assert output == (
        {"error": "provider_unavailable"} if network_probe else {"value": [1.0, -2.0]}
    )


@pytest.mark.parametrize("network_probe", [False, True])
def test_actual_child_kernel_limits_and_forbidden_attempt_count(network_probe):
    # Real isolated interpreter and real limits. Substitute only document parsing
    # after the fences, so the test can inspect them without weakening the child.
    script = """
import importlib.util, json, resource, socket
spec = importlib.util.spec_from_file_location('bounded_child', __import__('sys').argv[1])
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)
def parse(value):
    assert resource.getrlimit(resource.RLIMIT_AS) == (256 * 1024 * 1024,) * 2
    assert resource.getrlimit(resource.RLIMIT_CPU) == (15, 15)
    assert resource.getrlimit(resource.RLIMIT_CORE) == (0, 0)
    assert 'activekg.engine.embedding_provider' not in __import__('sys').modules
    if __import__('sys').argv[2] == 'probe':
        try:
            socket.getaddrinfo('not-a-real-destination.invalid', 443)
        except module.WorkRefused:
            pass
    return 'limits-proven'
module._parse_document = parse
module._run_child('parse')
"""
    result = subprocess.run(
        [
            sys.executable,
            "-I",
            "-B",
            "-c",
            script,
            providers.__file__,
            "probe" if network_probe else "clean",
        ],
        input=b'{"original_bytes":"synthetic"}',
        capture_output=True,
        timeout=5,
        env=providers._child_environment(),
    )
    assert result.returncode == 0 and not result.stderr
    assert json.loads(result.stdout) == (
        {"error": "provider_unavailable"} if network_probe else {"value": "limits-proven"}
    )


@pytest.mark.parametrize("tokens,expected", [(2, "value"), (99, "error")])
def test_embedding_child_never_silently_truncates_to_token_limit(tmp_path, tokens, expected):
    path = tmp_path / "snapshots" / ("a" * 40)
    path.mkdir(parents=True)
    scratch = tmp_path / "import"
    scratch.mkdir(mode=0o700)
    script = """
import importlib.util, sys, types
spec = importlib.util.spec_from_file_location('bounded_child', sys.argv[1])
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)
class Tokenizer:
    def encode(self, text, **kwargs):
        assert kwargs == {'truncation': False, 'add_special_tokens': True}
        return [1] * int(sys.argv[2])
class Fake:
    def __init__(self, **kwargs):
        assert kwargs['backend'] == 'sentence-transformers'
        self._model = types.SimpleNamespace(tokenizer=Tokenizer(), max_seq_length=10)
    def _ensure_model(self): pass
    def encode(self, texts):
        assert int(sys.argv[2]) <= 10
        return types.SimpleNamespace(tolist=lambda: [[1.] + [0.] * 383 for _ in texts])
stub = types.ModuleType('activekg.engine.embedding_provider')
stub.EmbeddingProvider = Fake
sys.modules[stub.__name__] = stub
sys.modules['torch'] = types.ModuleType('torch')
sys.modules['sentence_transformers'] = types.ModuleType('sentence_transformers')
module._run_child('embed')
"""
    payload = {"path": str(path), "texts": ["Python"], "revision": "a" * 40}
    result = subprocess.run(
        [sys.executable, "-I", "-B", "-c", script, providers.__file__, str(tokens)],
        input=json.dumps(payload).encode(),
        capture_output=True,
        timeout=5,
        env={**providers._child_environment(), "CANDIDATE_INDEX_IMPORT_DIR": str(scratch)},
    )
    assert result.returncode == 0 and not result.stderr
    output = json.loads(result.stdout)
    output.pop("fence", None)
    assert list(output) == [expected]
    if expected == "error":
        assert output["error"] == "chunk_overflow"
    else:
        assert output["value"] == [[1.0] + [0.0] * 383]


class Repository:
    tunables = IndexTunables()

    def __init__(self, *, attempt=1, reserve=True, stage="extract"):
        self.events = []
        self.should_reserve = reserve
        self.eligible = True
        self.completes = []
        self.failures = []
        self.dispatch = {
            "job_id": "job",
            "scope_key": "org_17",
            "source_id": "source",
            "generation_id": "generation",
            "source_kind": "organization_application",
            "stage": stage,
            "attempt": attempt,
            "policy": asdict(POLICY),
            "policy_sha256": POLICY.digest,
            "remaining_ms": 2000,
            "content_kind": "pinned_text",
            "professional_text": TEXT,
            "source_manifest": {"source_hash": sha256(TEXT)},
            "extraction": {"chunks": chunk_manifest(TEXT, POLICY)},
        }

    def reserve(self, lease, policy):
        self.events.append("reserve_commit")
        return self.dispatch if self.should_reserve else None

    def status(self, scope, source):
        self.events.append("privacy_commit")
        return {"eligible": self.eligible, "generation_id": "generation"}

    def complete_extract(self, dispatch, **kwargs):
        self.events.append("extract_commit")
        self.completes.append(kwargs)
        return self.eligible

    def complete_embed(self, dispatch, vectors):
        self.completes.append(vectors)
        return self.eligible

    def fail(self, dispatch, reason, retry_ms):
        self.failures.append((reason, retry_ms))
        return True


@pytest.mark.parametrize(
    "attempt,model", [(1, POLICY.primary_model_id), (2, POLICY.fallback_model_id)]
)
def test_worker_reserves_exact_attempt_then_fresh_privacy_then_one_http_then_commit(attempt, model):
    repo = Repository(attempt=attempt)

    def handler(request):
        repo.events.append("http")
        assert json.loads(request.content)["model"] == model
        return response(wire(model=model))

    worker = GenerationWorker(
        repo, POLICY, api_key="synthetic", transport=httpx.MockTransport(handler)
    )
    assert asyncio.run(worker.process({})) == "completed"
    assert repo.events == [
        "reserve_commit",
        "privacy_commit",
        "privacy_commit",
        "http",
        "extract_commit",
    ]
    assert repo.completes[0]["model_id"] == model and not repo.failures


def test_worker_reservation_refusal_zero_http_and_no_new_attempt():
    repo = Repository(reserve=False)
    worker = GenerationWorker(
        repo,
        POLICY,
        api_key="synthetic",
        transport=httpx.MockTransport(lambda _: pytest.fail("no dispatch allowed")),
    )
    assert asyncio.run(worker.process({})) == "not_reserved"
    assert repo.events == ["reserve_commit"] and not repo.failures


def test_worker_no_nested_retry_on_bad_primary_and_expiry_refuses_dispatch():
    repo = Repository()
    calls = []

    def handler(_):
        calls.append(1)
        return response(wire({**RESULT, "confidence": 0.4}))

    worker = GenerationWorker(
        repo, POLICY, api_key="synthetic", transport=httpx.MockTransport(handler)
    )
    assert asyncio.run(worker.process({})) == "low_confidence"
    assert calls == [1] and repo.failures == [("low_confidence", 0)] and not repo.completes
    repo.dispatch["remaining_ms"] = 0
    assert asyncio.run(worker.process({})) == "dispatch_exhausted"
    assert calls == [1]


def test_privacy_changed_during_model_work_prevents_completion():
    repo = Repository()

    def handler(_):
        repo.eligible = False
        return response(wire())

    worker = GenerationWorker(
        repo, POLICY, api_key="synthetic", transport=httpx.MockTransport(handler)
    )
    assert asyncio.run(worker.process({})) == "fenced"


def test_embedding_incomplete_vectors_never_publish(monkeypatch):
    repo = Repository(stage="embed")

    async def broken(*args, **kwargs):
        return [[1.0]]

    monkeypatch.setattr(processing, "embed_once", broken)
    assert asyncio.run(GenerationWorker(repo, POLICY).process({})) == "incomplete"
    assert not repo.completes


def test_parser_wait_is_followed_by_fresh_privacy_before_http(monkeypatch):
    repo = Repository()
    repo.dispatch.update(
        content_kind="original_bytes", professional_text=None, original_bytes="synthetic"
    )

    async def parser(*args, **kwargs):
        repo.events.append("parse")
        repo.eligible = False
        return TEXT

    monkeypatch.setattr(processing, "parse_original", parser)
    worker = GenerationWorker(
        repo,
        POLICY,
        api_key="synthetic",
        transport=httpx.MockTransport(lambda _: pytest.fail("privacy changed during parsing")),
    )
    assert asyncio.run(worker.process({})) == "privacy_restricted"
    assert repo.events == ["reserve_commit", "privacy_commit", "parse", "privacy_commit"]
    assert not repo.completes


def runtime_environment(monkeypatch, stage):
    from activekg.candidate_index import admission
    from activekg.extraction.schema import ExtractionResult

    for name in list(os.environ):
        if name.startswith(("CANDIDATE_INDEX_", "CANDIDATE_PRIVACY_HMAC_")):
            monkeypatch.delenv(name)
    policy = IndexPolicy(
        extraction_schema_sha256=sha256(canonical_json(ExtractionResult.model_json_schema())),
        extraction_prompt_sha256=sha256(
            (Path(admission.__file__).parents[1] / "extraction" / "prompt.py").read_bytes()
        ),
        primary_model_id=POLICY.primary_model_id,
        fallback_model_id=POLICY.fallback_model_id,
        embedding_model_id=POLICY.embedding_model_id,
        embedding_artifact_revision=POLICY.embedding_artifact_revision,
    )
    for name, value in {
        admission.FLAGS[stage]: "true",
        "EXTRACTION_PRIMARY_MODEL": policy.primary_model_id,
        "EXTRACTION_FALLBACK_MODEL": policy.fallback_model_id,
        "EMBEDDING_BACKEND": "sentence-transformers",
        "EMBEDDING_MODEL": policy.embedding_model_id,
        "CANDIDATE_INDEX_EMBEDDING_ARTIFACT_REVISION": policy.embedding_artifact_revision,
        "CANDIDATE_INDEX_POLICY_SHA256": policy.digest,
        "GROQ_API_KEY": "synthetic-key-not-used",
    }.items():
        monkeypatch.setenv(name, value)
    return policy


@pytest.mark.parametrize("stage", ["extract", "embed"])
@pytest.mark.parametrize("wrong", ["false", "true", ""])
def test_real_worker_entrypoint_refuses_misplaced_flag_before_schema_or_redis(
    monkeypatch, stage, wrong
):
    from activekg.common import metrics, schema_control
    from activekg.embedding import worker as embedding
    from activekg.extraction import worker as extraction

    runtime_environment(monkeypatch, stage)
    monkeypatch.setenv("CANDIDATE_INDEX_API_ENABLED", wrong)
    monkeypatch.setattr(
        schema_control, "assert_startup_schema_ready", lambda **_: pytest.fail("schema opened")
    )
    monkeypatch.setattr(metrics, "get_redis_client", lambda: pytest.fail("redis opened"))
    with pytest.raises(IndexContractError, match="flag_placement"):
        (extraction.start_extraction_worker if stage == "extract" else embedding.start_worker)()


@pytest.mark.parametrize("stage", ["extract", "embed"])
def test_worker_config_checks_pins_and_forbids_privacy_keyring(monkeypatch, stage):
    from activekg.privacy.config import CandidatePrivacyConfigurationError

    policy = runtime_environment(monkeypatch, stage)
    artifacts = []
    monkeypatch.setattr(
        providers, "cached_embedding_path", lambda *args: artifacts.append(args) or Path("/unused")
    )
    assert worker_configuration(stage).policy == policy
    assert artifacts == (
        [(policy.embedding_model_id, policy.embedding_artifact_revision)]
        if stage == "embed"
        else []
    )
    monkeypatch.setenv("CANDIDATE_PRIVACY_HMAC_ACTIVE_VERSION", "1")
    monkeypatch.setenv("CANDIDATE_PRIVACY_HMAC_KEY_V1", base64.b64encode(b"x" * 32).decode())
    with pytest.raises(CandidatePrivacyConfigurationError):
        worker_configuration(stage)
    monkeypatch.delenv("CANDIDATE_PRIVACY_HMAC_ACTIVE_VERSION")
    monkeypatch.delenv("CANDIDATE_PRIVACY_HMAC_KEY_V1")
    monkeypatch.setenv("CANDIDATE_INDEX_POLICY_SHA256", "0" * 64)
    with pytest.raises(IndexContractError, match="policy_not_pinned"):
        worker_configuration(stage)


def test_disabled_lane_needs_no_policy_or_artifact_and_enabled_prerequisites_refuse(monkeypatch):
    runtime_environment(monkeypatch, "extract")
    monkeypatch.delenv("GROQ_API_KEY")
    with pytest.raises(IndexContractError, match="extraction_key_missing"):
        worker_configuration("extract")
    monkeypatch.setenv("CANDIDATE_INDEX_EXTRACTION_ENABLED", "false")
    monkeypatch.delenv("CANDIDATE_INDEX_POLICY_SHA256")
    assert worker_configuration("extract") is None

    runtime_environment(monkeypatch, "embed")

    def missing(*_):
        raise WorkRefused("provider_unavailable")

    monkeypatch.setattr(providers, "cached_embedding_path", missing)
    with pytest.raises(IndexContractError, match="embedding_artifact_missing"):
        worker_configuration("embed")
    monkeypatch.setenv("CANDIDATE_INDEX_EMBEDDING_ENABLED", "false")
    assert worker_configuration("embed") is None


@pytest.mark.parametrize("stage", ["extract", "embed"])
@pytest.mark.parametrize("active", [False, True])
def test_actual_worker_entrypoint_starts_joins_lane_and_preserves_legacy_configuration(
    monkeypatch, stage, active
):
    from activekg.candidate_index import admission
    from activekg.common import metrics, schema_control
    from activekg.embedding import global_candidates
    from activekg.embedding import worker as embedding
    from activekg.extraction import worker as extraction

    policy = runtime_environment(monkeypatch, stage)
    monkeypatch.setenv(admission.FLAGS[stage], str(active).lower())
    monkeypatch.setenv("GLOBAL_MEMORY_ENABLED", "true")
    monkeypatch.setattr(providers, "cached_embedding_path", lambda *_: Path("/unused"))
    events = []
    seen = []
    ticked = threading.Event()

    def schema(**kwargs):
        assert kwargs == {"require_privacy_hmac": False}
        events.append("schema")
        return "synthetic-unused-dsn"

    def start_health(state):
        events.append("health")
        return SimpleNamespace(
            shutdown=lambda: events.append("health_stop"),
            server_close=lambda: events.append("health_close"),
        )

    async def tick(self, selected_stage):
        assert selected_stage == stage and self.policy == policy
        assert self.repository._dsn == "synthetic-unused-dsn"
        events.append("index_tick")
        ticked.set()
        return []

    def run(self):
        seen.append(self)
        events.append("legacy_run")
        if active:
            assert ticked.wait(3)
            assert self.index_runtime._thread.is_alive()
        else:
            assert self.index_runtime is None
        self._shutdown_handler(15, None)

    monkeypatch.setattr(schema_control, "assert_startup_schema_ready", schema)
    monkeypatch.setattr(metrics, "get_redis_client", lambda: events.append("redis") or object())
    monkeypatch.setattr(extraction, "start_healthcheck_server", start_health)
    monkeypatch.setattr(
        extraction,
        "start_database_monitor",
        lambda *_: SimpleNamespace(join=lambda **_: events.append("monitor_join")),
    )
    monkeypatch.setattr(extraction, "assert_extraction_models_configured", lambda: None)
    monkeypatch.setattr(extraction, "ExtractionClient", lambda **_: object())
    monkeypatch.setattr(extraction.signal, "signal", lambda *_: None)
    for module in (extraction, embedding):
        monkeypatch.setattr(module, "GraphRepository", lambda *_: object())
        monkeypatch.setattr(module, "CandidatePrivacyRepository", lambda *_: object())
    legacy_embedding = object()
    legacy_global = object()
    embed_calls = []
    global_calls = []
    monkeypatch.setattr(
        embedding, "EmbeddingProvider", lambda **kw: embed_calls.append(kw) or legacy_embedding
    )
    monkeypatch.setattr(
        global_candidates,
        "GlobalCandidateEmbedder",
        lambda *args, **kw: global_calls.append((args, kw)) or legacy_global,
    )
    monkeypatch.setattr(GenerationWorker, "tick", tick)
    monkeypatch.setattr(extraction.ExtractionWorker, "run", run)
    monkeypatch.setattr(embedding.EmbeddingWorker, "run", run)
    (extraction.start_extraction_worker if stage == "extract" else embedding.start_worker)()
    assert events.index("schema") < events.index("redis") < events.index("legacy_run")
    assert ("index_tick" in events) is active
    if stage == "embed":
        assert embed_calls == [
            {"backend": "sentence-transformers", "model_name": policy.embedding_model_id}
        ]
        assert global_calls[0][0] == ("synthetic-unused-dsn", legacy_embedding)
        assert global_calls[0][1] == {"batch_size": 64, "sweep_interval_seconds": 15.0}
        assert seen[0].global_candidate_embedder is legacy_global
    if active:
        assert seen[0].index_runtime.status() == "stopped"
        assert not seen[0].index_runtime._thread.is_alive()
        assert events[-3:] == ["health_stop", "health_close", "monitor_join"]


def test_runtime_close_cancels_and_reaps_real_child_without_duplicate_cancellation(monkeypatch):
    real = asyncio.create_subprocess_exec
    children = []
    child_ready = threading.Event()

    async def launch(*args, **kwargs):
        child = await real(args[0], "-I", "-c", "import time; time.sleep(20)", **kwargs)
        children.append(child)
        child_ready.set()
        return child

    async def tick(_stage):
        await providers.child_work(
            "parse", {"original_bytes": "YQ=="}, deadline=time.monotonic() + 15
        )
        pytest.fail("cancelled work must not complete")

    worker = GenerationWorker(Repository(), POLICY)
    monkeypatch.setattr(worker, "tick", tick)
    monkeypatch.setattr(asyncio, "create_subprocess_exec", launch)
    runtime = GenerationRuntime(worker, "extract")
    try:
        runtime.start()
        assert child_ready.wait(3)
        runtime.request_stop()
        runtime.request_stop()
    finally:
        runtime.close()
    assert runtime.status() == "stopped" and not runtime._thread.is_alive()
    assert children[0].returncode is not None
    with pytest.raises(ProcessLookupError):
        os.kill(children[0].pid, 0)
    with pytest.raises(IndexContractError, match="already_started"):
        runtime.start()


@pytest.mark.parametrize("stage", ["extract", "embed"])
def test_worker_readiness_is_index_aware_in_memory_and_has_no_false_ready(monkeypatch, stage):
    from activekg.extraction.worker import WorkerHealthState, worker_health_response

    runtime = GenerationRuntime(GenerationWorker(Repository(), POLICY), stage)
    state = WorkerHealthState(
        1,
        index_runtime=runtime,
        service=("extraction-worker" if stage == "extract" else "embedding-worker"),
    )
    state.provider_configured()
    state.loop_cycle_success()
    state.database_success()
    monkeypatch.setenv("ACTIVEKG_CONTROL_PLANE_TOKEN", "synthetic-control-token-long-enough")
    assert worker_health_response("/health", None, state)[0] == 200
    assert worker_health_response("/readyz", None, state)[0] == 401
    assert state.snapshot()[1]["candidate_index"] == "starting"
    assert not state.snapshot()[0]
    runtime._observe("ready")
    assert state.snapshot()[0]
    runtime._observe("error")
    assert not state.snapshot()[0]
    runtime._observe("ready")
    runtime._observed_at -= 180
    assert state.snapshot()[1]["candidate_index"] == "stale"
    assert not state.snapshot()[0]
    runtime.close()
    assert state.snapshot()[1]["candidate_index"] == "stopped"
